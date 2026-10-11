"""사건 작업 공간(빠진 정보·경위·증거·표현 바꾸기·함께 내려받기)과 캘린더 파일 검사.

순수 함수는 바로, 화면은 가짜 Supabase·Gemini로 실제 app.py를 실행해 확인한다. 실행: python3 -m unittest discover -s tests
"""

import datetime
import os
import re
import sys
import unittest
from io import BytesIO
from unittest import mock

from docx import Document
from streamlit.testing.v1 import AppTest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import calendar_ics as cal  # noqa: E402
import case_workspace as w  # noqa: E402
import fake_services as fake  # noqa: E402

CASE = {
    "case_id": "44444444-4444-4444-8444-444444444444", "user_id": fake.ADMIN_ID, "case_type": "중고거래 사기",
    "reporter_info": "홍길동(시험)", "incident_datetime": "9월 3일",
    "incident_description": "9월 2일 오후 3시 ○○마켓에서 판매글을 보고 45만 원 송금 → 판매자가 내일 보낸다고 함 → 9월 4일부터 연락 끊김",
    "damage_details": "50만 원 송금", "evidence": "채팅 캡처, 이체 내역", "suspect_info": "모름",
    "created_at": "2026-10-01T00:00:00+00:00",
}


class RulesTest(unittest.TestCase):
    def test_missing_questions_skip_unknown_and_keep_order(self):
        keys = [k for k, _, _ in w.missing_questions(CASE)]
        self.assertEqual(keys, ["incident_location", "witness", "punishment_intent"])
        self.assertEqual(w.unknown_items(CASE), ["상대방 정보"])

    def test_inconsistent_amounts_and_dates_are_asked_not_judged(self):
        notes = w.inconsistencies(CASE)
        self.assertEqual(len(notes), 2)
        self.assertIn("45만 원", notes[0])
        self.assertIn("50만 원", notes[0])
        self.assertIn("확인해 주세요", notes[1])
        self.assertEqual(w.inconsistencies({**CASE, "damage_details": "45만 원", "incident_datetime": "9월 2일"}), [])

    def test_timeline_only_copies_written_dates(self):
        rows = w.build_timeline(CASE)
        self.assertEqual([r["when"] for r in rows], ["9월 2일 오후 3시", "", "9월 4일"])
        joined = " ".join(r["what"] for r in rows)
        story_digits = set(re.findall(r"\d+", CASE["incident_description"]))
        self.assertTrue(set(re.findall(r"\d+", joined + " ".join(r["when"] for r in rows))) <= story_digits)

    def test_evidence_and_links(self):
        items = w.evidence_items(CASE, ["capture.jpg"])
        self.assertEqual([i["name"] for i in items], ["채팅 캡처", "이체 내역", "capture.jpg"])
        self.assertTrue(all(i["shows"] == "" for i in items))  # 무엇을 보여 주는지는 사용자가 적는다
        rows = w.build_timeline(CASE)
        rows[0]["evidence"] = "이체 내역"
        issues = w.link_issues(rows, items)
        self.assertEqual(len(issues["events_without_evidence"]), 2)
        self.assertEqual(issues["evidence_not_used"], ["채팅 캡처", "capture.jpg"])

    def test_next_action_order(self):
        self.assertEqual(w.next_action(CASE, False, {}), "빠진 정보 확인하기")
        done = {**CASE, "incident_location": "○○마켓", "witness": "없음", "punishment_intent": "원함",
                "damage_details": "45만 원", "incident_datetime": "9월 2일"}
        self.assertEqual(w.next_action(done, False, {}), "초안 만들기")
        self.assertEqual(w.next_action(done, True, {}), "초안 내용 직접 확인하기")
        self.assertEqual(w.next_action(done, True, {"review": True}), "제출할지 정하고 표시하기")

    def test_rewrite_prompt_and_changes(self):
        prompt = w.rewrite_prompt("초안", "concise")
        for phrase in ["표현만", "새로 넣거나 빼지 마세요", "[확인 필요]", "강도를 바꾸지"]:
            self.assertIn(phrase, prompt)
        self.assertEqual(w.rewrite_changes("9월 2일 45만 원 [확인 필요]", "9월 2일에 45만 원을 [확인 필요]"), [])
        notes = w.rewrite_changes("9월 2일 45만 원 [확인 필요]", "9월 3일 45만 원")
        self.assertEqual(len(notes), 3)

    def test_bundle_includes_only_chosen_parts_and_notice(self):
        rows, items = w.build_timeline(CASE), w.evidence_items(CASE)
        text = w.bundle_text("고소장\n본문", rows, items, {"draft": False, "timeline": True, "evidence": True})
        self.assertNotIn("고소장\n본문", text)
        self.assertIn("날짜 확인 필요", text)
        self.assertIn("파일 내용은 고소미가 확인하지 않았습니다", text)
        self.assertIn("법률 전문가의 검토", text)
        doc = Document(BytesIO(w.bundle_docx(text).getvalue()))
        self.assertIn("사건 경위 (직접 정리한 순서)", [p.text for p in doc.paragraphs])


class CalendarTest(unittest.TestCase):
    def test_ics_dates_times_alarm_and_escaping(self):
        text = cal.build_ics([
            {"kind": "submit", "title": "경찰서, 제출; 준비", "date": datetime.date(2026, 10, 20),
             "time": datetime.time(14, 30), "alarm": 60, "memo": "신분증\n자료"},
            {"kind": "consult", "title": "", "date": datetime.date(2026, 10, 31), "time": None, "alarm": 0},
        ], "중고거래 사기", now=datetime.datetime(2026, 10, 11, 0, 0, tzinfo=datetime.timezone.utc))
        self.assertTrue(text.startswith("BEGIN:VCALENDAR\r\n") and text.endswith("END:VCALENDAR\r\n"))
        unfolded = text.replace("\r\n ", "")
        self.assertIn("DTSTART;TZID=Asia/Seoul:20261020T143000", unfolded)
        self.assertIn("DTEND;TZID=Asia/Seoul:20261020T153000", unfolded)
        self.assertIn("TRIGGER:-PT60M", unfolded)
        self.assertIn("SUMMARY:[고소미] 경찰서\\, 제출\; 준비 · 중고거래 사기", unfolded)
        self.assertIn("신분증\\n자료", unfolded)
        self.assertIn("DTSTART;VALUE=DATE:20261031", unfolded)
        self.assertIn("DTEND;VALUE=DATE:20261101", unfolded)
        self.assertIn("SUMMARY:[고소미] 상담", unfolded)
        self.assertEqual(unfolded.count("BEGIN:VALARM"), 1)
        self.assertIn("DTSTAMP:20261011T000000Z", unfolded)
        for line in text.split("\r\n"):
            self.assertLessEqual(len(line.encode("utf-8")), 75)

    def test_rejects_missing_date_and_bad_alarm(self):
        with self.assertRaises(ValueError):
            cal.clean_event({"title": "x", "date": None})
        with self.assertRaises(ValueError):
            cal.clean_event({"title": "x", "date": datetime.date(2026, 1, 1), "alarm": 7})


ENV = {
    "SUPABASE_URL": "https://fake.invalid", "SUPABASE_PUBLISHABLE_KEY": "fake", "GEMINI_API_KEY": "fake",
    "ADMIN_USER_IDS": fake.ADMIN_ID, "PRIVACY_POLICY_ENABLED": "true", "PRIVACY_OPERATOR_NAME": "운영자",
    "PRIVACY_CONTACT_EMAIL": "contact@example.com", "PRIVACY_POLICY_EFFECTIVE_DATE": "2026-09-27", "AI_CALLS_PER_HOUR": "30",
}
CID = CASE["case_id"]


class WorkspaceAppTest(unittest.TestCase):
    def setUp(self):
        fake.reset()
        fake.CASES.append(dict(CASE))
        patches = [mock.patch.dict(os.environ, ENV), mock.patch("supabase.create_client", fake.fake_create_client),
                   mock.patch("google.genai.Client", fake.FakeGenaiClient)]
        for patch in patches:
            patch.start()
        self.addCleanup(lambda: [patch.stop() for patch in reversed(patches)])

    def open(self, user_id=fake.ADMIN_ID, step=None):
        at = AppTest.from_file(os.path.join(ROOT, "app.py"), default_timeout=60)
        at.session_state["supabase_session"] = fake.session_for(user_id, "x@example.com")
        at.session_state["browser_login_checked"] = True
        at.run()
        at.session_state["nav_page"] = "cases"
        at.session_state["open_case_id"] = CID
        if step:
            at.session_state[f"ws_step_{CID}"] = step
        at.run()
        self.assertEqual([e.value for e in at.exception], [])
        return at

    def test_regular_user_keeps_old_case_detail(self):
        fake.CASES[0]["user_id"] = fake.USER_ID
        at = self.open(fake.USER_ID)
        self.assertNotIn(f"ws_{CID}", at.session_state)
        self.assertIn("고소장 초안", [e.label for e in at.expander])

    def test_check_step_asks_one_question_and_accepts_unknown(self):
        at = self.open()
        self.assertTrue(any("다음 할 일: 빠진 정보 확인하기" in m.value for m in at.markdown))
        self.assertEqual(len([w_ for w_ in at.warning if "확인해 주세요" in w_.value]), 2)
        self.assertTrue(any("어디에서 있었던 일인가요?" in m.value for m in at.markdown))
        at.button(key=f"ws_unknown_{CID}_incident_location").click().run()
        self.assertEqual(at.session_state[f"ws_{CID}"]["answers"]["incident_location"], "모름")
        self.assertTrue(any("본 사람이 있나요?" in m.value for m in at.markdown))
        self.assertFalse([w_ for w_ in fake.WRITES if w_[0] == "cases"])  # 저장된 사건은 바꾸지 않는다

    def test_story_step_builds_timeline_and_evidence(self):
        at = self.open(step="story")
        ws = at.session_state[f"ws_{CID}"]
        self.assertEqual([r["when"] for r in ws["timeline"]], ["9월 2일 오후 3시", "", "9월 4일"])
        self.assertEqual([e["name"] for e in ws["evidence"]], ["채팅 캡처", "이체 내역"])

    def test_rewrite_warns_on_changed_numbers_and_undo_restores(self):
        fake.REPLIES[:] = ["고소장\n범죄 사실\n9월 2일 45만 원 송금 [확인 필요]", "고소장\n범죄 사실\n9월 3일 45만 원 송금"]
        at = self.open(step="draft")
        at.button(key=f"ws_make_{CID}").click().run()
        before = at.text_area(key=f"list_draft_editor_{CID}").value
        at.button(key=f"ws_style_{CID}_concise").click().run()
        self.assertTrue(any("숫자" in w_.value for w_ in at.warning))
        self.assertIn("9월 3일", at.text_area(key=f"list_draft_editor_{CID}").value)
        at.button(key=f"ws_undo_{CID}").click().run()
        self.assertEqual(at.text_area(key=f"list_draft_editor_{CID}").value, before)
        self.assertIn("표현만", fake.contents_text(fake.CALLS[-1]))

    def test_draft_ai_error_keeps_work(self):
        fake.REPLIES[:] = [RuntimeError("503")]
        at = self.open(step="draft")
        at.button(key=f"ws_make_{CID}").click().run()
        self.assertTrue(at.error)
        self.assertTrue(any(b.key == f"ws_make_{CID}" for b in at.button))

    def test_schedule_and_ics_download(self):
        at = self.open(step="next")
        at.button(key=f"ws_add_{CID}_schedules").click().run()
        p = f"ws_f_{CID}_schedules_"
        at.text_input(key=p + "title").input("경찰서 제출").run()
        at.date_input(key=p + "date").set_value(datetime.date(2026, 10, 20)).run()
        at.button(key=f"ws_fok_{CID}_schedules").click().run()
        events = at.session_state[f"ws_{CID}"]["schedules"]
        self.assertEqual(events[0]["date"], datetime.date(2026, 10, 20))
        self.assertIsNone(events[0]["time"])
        self.assertTrue(any("일정 앱이 보내" in c.value or "일정 앱에서" in c.value for c in at.caption))

    def test_review_asks_missing_info_before_saving_and_updates_input(self):
        at = AppTest.from_file(os.path.join(ROOT, "app.py"), default_timeout=60)
        at.session_state["supabase_session"] = fake.session_for(fake.ADMIN_ID, "x@example.com")
        at.session_state["browser_login_checked"] = True
        at.run()
        at.session_state["nav_page"] = "case"
        at.run()
        self.assertTrue(any(b.key == "case_start_text" for b in at.button))  # 글로 이야기하기로 시작 가능
        at.selectbox(key="case_type_choice").select("중고거래 사기").run()
        at.text_area(key="case_field_reporter_info").input("홍길동").run()
        at.text_area(key="case_field_incident_description").input("9월 2일 45만 원 송금").run()
        at.checkbox(key="case_data_consent").check().run()
        at.button(key="case_form_submit").click().run()
        self.assertTrue(any("언제 있었던 일인가요?" in m.value for m in at.markdown))
        at.text_input(key="review_ans_incident_datetime").input("9월 2일 오후 3시").run()
        at.button(key="review_put_incident_datetime").click().run()
        self.assertEqual(at.session_state["case_field_incident_datetime"], "9월 2일 오후 3시")
        self.assertEqual(at.session_state["case_review"]["values"]["incident_datetime"], "9월 2일 오후 3시")
        at.button(key="review_skip_suspect_info").click().run()
        self.assertTrue(any("어디에서 있었던 일인가요?" in m.value for m in at.markdown))
        self.assertFalse([w_ for w_ in fake.WRITES if w_[0] == "cases"])  # 확인 단계에서는 저장하지 않는다


if __name__ == "__main__":
    unittest.main()
