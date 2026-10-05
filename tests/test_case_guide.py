"""case_guide.py 검사. 실행: python3 -m unittest discover -s tests"""

import os
import sys
import unittest
from urllib.parse import urlparse

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import case_guide as g  # noqa: E402

OFFICIAL_HOST_SUFFIXES = (".go.kr", ".or.kr")


def filled_case(**overrides):
    case = {
        "case_type": "중고거래 사기",
        "reporter_info": "홍길동 / 010-1111-2222",
        "suspect_info": "이름 모름 / 닉네임 판매왕 / 계좌 123-456-7890",
        "incident_datetime": "2026년 9월 2일 오후 3시",
        "incident_location": "온라인(○○마켓 앱 채팅)",
        "incident_description": "9월 2일 태블릿 판매글을 보고 45만 원을 송금했고 9월 4일부터 연락이 끊겼습니다.",
        "damage_details": "45만 원",
        "evidence": "채팅 캡처, 이체 내역",
        "witness": "없음",
        "punishment_intent": "처벌을 원합니다",
    }
    case.update(overrides)
    return case


class DamageTypeDataTest(unittest.TestCase):
    def test_required_types_are_present(self):
        for name in [
            "중고거래 사기", "온라인 쇼핑 사기", "투자 사기", "명예훼손·모욕", "협박·폭행", "스토킹",
            "성범죄·디지털 성범죄", "학교폭력", "직장 내 괴롭힘", "개인정보 침해·명의도용",
        ]:
            self.assertIn(name, g.DAMAGE_TYPES)

    def test_every_type_is_complete_and_cites_official_sources(self):
        for name, guide in g.DAMAGE_TYPES.items():
            with self.subTest(name=name):
                self.assertIn(guide["route"], g.ROUTES)
                self.assertTrue(guide["summary"])
                self.assertTrue(guide["checklist"])
                self.assertTrue(guide["evidence_tips"])
                self.assertTrue(guide["qa"])
                self.assertTrue(guide["examples"].get("incident_description", "").startswith("예:"))
                self.assertTrue(guide["sources"])
                for _, url in guide["sources"]:
                    host = urlparse(url).hostname or ""
                    self.assertEqual(urlparse(url).scheme, "https")
                    self.assertTrue(host.endswith(OFFICIAL_HOST_SUFFIXES), url)

    def test_non_criminal_routes_are_marked(self):
        self.assertEqual(g.DAMAGE_TYPES["임금·퇴직금 미지급"]["route"], "labor")
        self.assertEqual(g.DAMAGE_TYPES["빌려준 돈·계약 분쟁"]["route"], "civil")
        self.assertEqual(g.DAMAGE_TYPES["직장 내 괴롭힘"]["route"], "mixed")

    def test_guide_text_has_no_statute_numbers_or_deadlines(self):
        # 확인되지 않은 법조문·기한을 안내 문구에 넣지 않는다.
        for name, guide in g.DAMAGE_TYPES.items():
            text = " ".join(
                [guide["summary"], guide.get("route_detail", ""), guide.get("urgent", "")]
                + [q + a for q, a in guide["qa"]]
            )
            with self.subTest(name=name):
                self.assertNotRegex(text, r"제\s*\d+\s*조")
                self.assertNotRegex(text, r"\d+\s*(개월|년)\s*(이내|안에)")

    def test_placeholder_falls_back_to_generic_guide(self):
        self.assertIn("판매글", g.field_placeholder("incident_description", "중고거래 사기"))
        self.assertEqual(g.field_placeholder("witness", "중고거래 사기"), g.FIELD_GUIDES["witness"]["placeholder"])
        self.assertEqual(
            g.field_placeholder("incident_description", g.OTHER_TYPE),
            g.FIELD_GUIDES["incident_description"]["placeholder"],
        )

    def test_every_field_except_type_has_a_guide(self):
        for _, key in g.CASE_FIELDS:
            if key != "case_type":
                self.assertIn(key, g.FIELD_GUIDES)


class CaseValueTest(unittest.TestCase):
    def test_required_fields(self):
        errors = g.validate_case_values({"case_type": "협박·폭행"})
        self.assertEqual(len(errors), 1)
        self.assertIn("신고인 정보", errors[0])
        self.assertIn("구체적인 사건 경위", errors[0])
        self.assertEqual(g.validate_case_values(filled_case()), [])

    def test_length_limit(self):
        errors = g.validate_case_values(filled_case(witness="가" * 1001))
        self.assertTrue(any("목격자 유무" in e for e in errors))

    def test_blank_optional_fields_become_missing_marker(self):
        values = g.normalize_case_values(filled_case(witness="  ", evidence=""))
        self.assertEqual(values["witness"], g.MISSING_VALUE)
        self.assertEqual(values["evidence"], g.MISSING_VALUE)
        self.assertEqual(values["case_type"], "중고거래 사기")

    def test_missing_info_items(self):
        items = g.missing_info_items(g.normalize_case_values(filled_case(witness="")))
        self.assertIn("목격자 유무 (입력하지 않음)", items)
        self.assertIn("상대방 정보 (모르는 부분 있음)", items)
        self.assertEqual(len(items), 2)

    def test_fingerprint_ignores_surrounding_spaces(self):
        a = g.compute_case_fingerprint("u1", filled_case())
        b = g.compute_case_fingerprint("u1", filled_case(witness=" 없음 "))
        c = g.compute_case_fingerprint("u2", filled_case())
        self.assertEqual(a, b)
        self.assertNotEqual(a, c)


class DraftPromptTest(unittest.TestCase):
    def test_prompt_includes_only_given_facts_and_safety_rules(self):
        prompt = g.build_draft_prompt(g.normalize_case_values(filled_case(witness="")), "2026-10-06")
        self.assertIn("- 목격자 유무: (입력 안 됨)", prompt)
        self.assertIn("- 피해 내용: 45만 원", prompt)
        self.assertIn("[확인 필요]", prompt)
        self.assertIn("성명불상", prompt)
        self.assertIn("법조문", prompt)
        self.assertIn("2026-10-06", prompt)
        self.assertNotIn("[참고]", prompt)

    def test_civil_type_adds_other_procedure_note(self):
        prompt = g.build_draft_prompt(filled_case(case_type="빌려준 돈·계약 분쟁"), "2026-10-06")
        self.assertIn("[참고]", prompt)


class UnsupportedDetailTest(unittest.TestCase):
    def test_faithful_draft_has_no_warnings(self):
        draft = (
            "고소장\n고소인\n홍길동 (010-1111-2222)\n피고소인\n성명불상 (닉네임 판매왕, 계좌 123-456-7890)\n"
            "범죄 사실\n고소인은 2026년 9월 2일 45만 원을 송금하였으나 9월 4일부터 연락이 되지 않았습니다.\n"
            "작성일\n2026-10-06"
        )
        self.assertEqual(g.find_unsupported_details(draft, filled_case(), "2026-10-06"), [])

    def test_invented_details_are_flagged(self):
        draft = (
            "피고소인은 형법 제347조의 사기죄를 범하였습니다. 9월 10일 추가로 100만 원을 받았고 "
            "연락처는 010-9999-8888입니다."
        )
        warnings = g.find_unsupported_details(draft, filled_case(), "2026-10-06")
        joined = "\n".join(warnings)
        self.assertIn("제347조", joined)
        self.assertIn("사기죄", joined)
        self.assertIn("100만원", joined)
        self.assertIn("9월 10일", joined)
        self.assertIn("010-9999-8888", joined)

    def test_iso_date_of_an_entered_date_is_not_flagged_as_a_number(self):
        draft = "고소인은 2026-09-02에 45만 원을 송금하였습니다."
        self.assertEqual(g.find_unsupported_details(draft, filled_case(), "2026-10-06"), [])

    def test_details_the_user_wrote_are_allowed(self):
        case = filled_case(incident_description="상대방이 형법 제347조 사기죄라고 스스로 말했습니다. 45만 원 송금.")
        draft = "고소인은 상대방이 '형법 제347조 사기죄'라고 말했다고 주장합니다."
        self.assertEqual(g.find_unsupported_details(draft, case, "2026-10-06"), [])


class DocxTest(unittest.TestCase):
    def paragraphs(self, text):
        document = Document(g.build_case_draft_docx(text))
        return [(p.text, any(run.bold for run in p.runs)) for p in document.paragraphs if p.text]

    def test_markdown_headings_are_cleaned_and_title_not_duplicated(self):
        text = "고소장\n## 고소인\n**홍길동**\n**피고소인:**\n성명불상\n1. 고소 취지\n수사를 요청합니다.\n" + g.DRAFT_DISCLAIMER
        paragraphs = self.paragraphs(text)
        texts = [t for t, _ in paragraphs]
        self.assertEqual(texts.count("고소장"), 1)
        self.assertIn(("고소인", True), paragraphs)
        self.assertIn(("피고소인", True), paragraphs)
        self.assertIn(("고소 취지", True), paragraphs)
        self.assertIn("홍길동", texts)
        self.assertIn(g.DRAFT_DISCLAIMER, texts)

    def test_edited_text_is_what_gets_exported(self):
        texts = [t for t, _ in self.paragraphs("고소인\n직접 고친 문장입니다.")]
        self.assertIn("직접 고친 문장입니다.", texts)


if __name__ == "__main__":
    unittest.main()
