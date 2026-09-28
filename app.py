import os
import base64
import json
import re
import time
import datetime
import unicodedata
import uuid
import hashlib
import html
import secrets
import threading
from io import BytesIO
from urllib.parse import urlencode
import streamlit as st
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from google import genai
from google.genai import types
from supabase import create_client, Client
from supabase_auth.errors import AuthApiError
from calendar_test import get_calendar_service
MEMORY_FILE = "memory.json"
MEMORY_TRIGGERS = [unicodedata.normalize("NFC", t) for t in ["기억해", "저장해"]]

SCHEDULE_FILE = "schedules.json"
CANCEL_TRIGGERS = [unicodedata.normalize("NFC", t) for t in ["취소", "삭제"]]
DATE_WORDS = ["오늘", "내일", "모레", "이번주", "이번 주", "다음주", "다음 주"]

CASES_TABLE = "cases"
CASE_FIELDS = [
    ("사건 유형", "case_type"),
    ("신고인 정보", "reporter_info"),
    ("상대방 정보", "suspect_info"),
    ("사건 발생 일시", "incident_datetime"),
    ("사건 발생 장소", "incident_location"),
    ("구체적인 사건 경위", "incident_description"),
    ("피해 내용", "damage_details"),
    ("증거 자료", "evidence"),
    ("목격자 유무", "witness"),
    ("처벌 의사", "punishment_intent"),
]
CASE_LONG_TEXT_FIELDS = {"reporter_info", "suspect_info", "incident_description", "damage_details", "evidence"}
CASE_FIELD_MAX_CHARS = {
    "case_type": 100,
    "reporter_info": 1000,
    "suspect_info": 1000,
    "incident_datetime": 100,
    "incident_location": 300,
    "incident_description": 8000,
    "damage_details": 4000,
    "evidence": 4000,
    "witness": 1000,
    "punishment_intent": 500,
}
DRAFT_DISCLAIMER = "이 문서는 참고용 초안이며 제출 전 법률 전문가의 검토가 필요합니다."

# 쉬운 설명은 형사 절차의 일반적인 뜻만 안내합니다. 사건별 판단은 하지 않습니다.
LEGAL_TERMS = [
    ("고소", "범죄 피해자 등 고소할 수 있는 사람이 수사기관에 범죄 사실을 알리고 처벌을 요청하는 것입니다."),
    ("고발", "고소할 수 있는 사람 외의 제3자가 수사기관에 범죄 사실을 알리고 처벌을 요청하는 것입니다."),
    ("피의자", "범죄 혐의를 받아 수사 대상이 된 사람입니다. 혐의가 확정됐다는 뜻은 아닙니다."),
    ("피고인", "검사가 기소해 형사재판을 받는 사람입니다. 유죄가 확정됐다는 뜻은 아닙니다."),
    ("송치", "수사기관이 사건 기록 등을 검찰로 넘기는 절차입니다. 송치 자체가 기소나 유죄를 뜻하지는 않습니다."),
    ("기소", "검사가 형사재판을 열어 달라고 법원에 공소를 제기하는 것입니다."),
    ("불기소", "검사가 해당 사건을 형사재판에 넘기지 않기로 결정하는 것입니다. 이유는 사건마다 다를 수 있습니다."),
]
LEGAL_TERMS_SOURCE = "https://www.spo.go.kr/site/spo/01/10101050200002018112210.jsp"

# 질문과 답변은 일반적인 정보와 기록 방법만 다루며 사건의 성립 여부를 판단하지 않습니다.
VICTIM_QA = {
    "사기 피해": {
        "source": "https://ecrm.police.go.kr/minwon/main",
        "questions": [
            ("돈을 보냈는데 상대방이 연락을 피합니다. 무엇부터 할까요?", "**먼저 거래 내용을 한 장에 정리해 보세요.** 무엇을 사거나 받기로 했는지, 언제 얼마를 어느 계좌로 보냈는지, 그 뒤 상대방이 어떤 말을 했는지 순서대로 적으시면 됩니다.\n\n예: ‘9월 2일 중고 거래 게시글을 보고 10만 원 송금 → 9월 3일 발송하겠다는 답변 → 9월 5일부터 연락이 닿지 않음.’ 연락이 끊겼다는 사실과 ‘처음부터 속일 생각이었다’는 추측은 구분해 주세요."),
            ("신고하려면 어떤 자료를 챙겨야 하나요?", "**거래가 어떻게 시작됐고 돈이 어디로 갔는지 보여 주는 자료**를 모아 두세요. 거래 글·상품 사진·상대방 계정이나 연락처, 메신저 대화 전체, 계좌이체 또는 결제 내역이 도움이 됩니다.\n\n화면을 캡처할 때는 날짜와 상대방 표시가 함께 보이게 남기고, 원본 대화나 파일도 삭제하지 마세요. 자료가 일부 없어도 기억나는 내용을 시간 순서대로 적을 수 있습니다."),
            ("인터넷으로 신고하면 경찰서 방문은 끝인가요?", "**온라인 접수 뒤에도 방문 요청을 받을 수 있습니다.** 경찰의 사이버범죄 신고시스템(ECRM)에서 신고할 수 있고, 가까운 경찰서에 직접 방문해 상담·접수할 수도 있습니다.\n\n온라인 신고를 하실 때는 피해 경위, 이체 내역, 대화 자료를 준비해 주세요. 접수 뒤에는 담당 경찰관이 추가 자료나 출석을 요청하는지 확인하시면 됩니다."),
        ],
    },
    "온라인 명예훼손": {
        "source": "https://www.easylaw.go.kr/CSP/CnpClsMain.laf?ccfNo=3&cciNo=1&cnpClsNo=1&csmSeq=293",
        "questions": [
            ("나에 관한 글을 발견했습니다. 삭제되기 전에 무엇을 해야 하나요?", "**글과 주변 상황을 함께 남겨 주세요.** 게시글 전체 화면, 주소(URL), 작성 날짜, 작성자 표시, 댓글을 캡처하고 가능하면 게시글이 올라온 곳도 기록하세요.\n\n글에서 정확히 어떤 표현이 쓰였는지, 다른 사람들이 볼 수 있었는지도 나중에 확인할 수 있게 남겨 두는 게 좋습니다. 캡처만 하고 원본 주소를 빼놓기 쉬우니 주소도 따로 저장해 주세요."),
            ("기분 나쁜 글이면 모두 신고할 수 있나요?", "**불쾌한 표현이라는 이유만으로 결과가 정해지지는 않습니다.** 글의 정확한 내용, 누구에 관한 글인지, 어디에 공개됐는지 등 상황을 함께 살펴봐야 합니다.\n\n우선 글의 원문과 주소를 보관하고, 사실로 확인한 내용과 본인의 생각을 나누어 적어 보세요. 고소음이는 해당 글의 위법 여부를 확정하지 않으므로, 판단이 필요하면 원문을 가지고 전문 기관이나 법률 전문가와 상담해 주세요."),
            ("상담할 때는 사건을 어떻게 설명하면 쉬울까요?", "**‘언제·어디에·어떤 글이 올라왔는지’부터 말씀해 보세요.** 그다음 글을 알게 된 경위, 본인과 글의 관련성, 글이 퍼진 상황을 순서대로 정리하시면 됩니다.\n\n예: ‘9월 4일 공개 게시판에서 제 이름이 적힌 글을 봤고, 댓글 3개가 달려 있었습니다.’ 확인하지 못한 작성자의 의도나 피해 규모를 단정해 적을 필요는 없습니다."),
        ],
    },
    "스토킹 피해": {
        "source": "https://www.easylaw.go.kr/CSP/CnpClsMain.laf?ccfNo=2&cciNo=2&cnpClsNo=1&csmSeq=1661",
        "questions": [
            ("상대방이 지금 집 앞에 있거나 따라옵니다. 어떻게 해야 하나요?", "**지금 위험하다면 안전한 곳으로 이동하고 112에 신고하세요.** 가능하면 주변 사람에게 상황을 알리고 혼자 상대방을 만나 해결하려고 하지 마세요.\n\n신고할 때는 현재 위치, 상대방이 있는 곳, 방금 일어난 일, 즉시 도움이 필요한 이유를 차례대로 말하면 됩니다. 고소음이에 내용을 적는 것은 경찰 신고로 연결되지 않습니다."),
            ("연락과 접근이 여러 번 있었는데 어떻게 정리하나요?", "**한 번씩 따로 기록해 보세요.** 날짜·시간·장소·어떤 행동이 있었는지, 당시 주변에 누가 있었는지를 시간 순서대로 적으시면 됩니다. 문자와 통화 기록, 이미 찍어 둔 사진이나 영상도 원본 그대로 보관하세요.\n\n예: ‘9월 5일 오후 8시, 집 앞에서 기다림 → 9월 6일 오전 7시, 전화 5회.’ 자료를 얻으려고 상대방을 찾아가거나 위험한 장소에 머무르지 마세요."),
            ("다시 찾아오지 못하게 도움을 요청할 수 있나요?", "**경찰에 현재 상황과 안전에 대한 걱정을 구체적으로 말해 주세요.** 반복된 연락이나 접근 내용, 다시 찾아올까 두려운 이유를 알리고 어떤 보호조치를 요청할 수 있는지 물어보시면 됩니다.\n\n접근이나 연락을 제한하는 조치가 검토될 수 있지만, 실제 적용 여부는 사건 상황과 담당 기관의 판단에 따라 달라집니다. 당장 위험하면 먼저 112에 신고하세요."),
        ],
    },
    "학교폭력 피해": {
        "source": "https://www.easylaw.go.kr/CSP/CnpClsMain.laf?ccfNo=3&cciNo=1&cnpClsNo=1&csmSeq=1408",
        "questions": [
            ("학교에서 괴롭힘을 당하고 있습니다. 누구에게 먼저 말하나요?", "**혼자 해결하려고 버티지 마시고 믿을 수 있는 어른에게 알려 주세요.** 담임교사나 학교폭력 담당 교사에게 언제 어디서 어떤 일이 있었는지 말씀할 수 있습니다. 보호자와 함께 이야기해도 됩니다.\n\n학교에 바로 말하기 어렵다면 **117 학교폭력 신고·상담센터**에 상담하거나 신고할 수 있습니다. 지금 폭행이나 위협을 받고 있다면 안전한 곳으로 이동하고 112에 신고하세요."),
            ("말로 설명하기 어려운데 무엇을 적어 가면 될까요?", "**기억나는 사건 하나씩 적어 가면 됩니다.** 날짜와 장소, 누가 무엇을 했는지, 곁에서 본 사람이 있었는지를 짧게 써 보세요.\n\n메신저 대화, 사진·영상, 이미 받은 진료 기록 등이 있다면 함께 보관하세요. 정확한 날짜가 생각나지 않으면 ‘9월 첫째 주쯤’처럼 기억나는 범위만 쓰고, 모르는 내용을 억지로 채우지 마세요."),
            ("알린 뒤에도 계속 마주쳐야 해서 걱정됩니다.", "**신고하거나 상담할 때 그 걱정을 함께 이야기해 주세요.** ‘복도에서 매일 마주친다’, ‘다시 연락이 올까 두렵다’처럼 지금 불안한 상황을 구체적으로 설명하면 담당자가 필요한 보호 방안을 함께 살펴볼 수 있습니다.\n\n학교에 알린 날짜와 받은 안내도 따로 적어 두세요. 당장 위험한 상황이라면 보호 방안 안내를 기다리지 말고 112에 신고하세요."),
        ],
    },
    "보이스피싱·문자 사기": {
        "source": "https://www.police.go.kr/www/security/cyber/cyber01.jsp",
        "questions": [
            ("방금 돈을 보냈습니다. 무엇부터 해야 하나요?", "**즉시 112에 전화해 보이스피싱 피해 사실을 알리세요.** 송금한 은행에도 바로 연락해 지급정지 등 가능한 조치를 문의하세요. 시간이 중요한 상황이므로 고소음이에 기록을 남기는 일보다 먼저 연락하는 게 좋습니다.\n\n전화할 때 송금 시각, 금액, 보낸 계좌와 받는 계좌를 확인할 수 있으면 도움이 됩니다. 모르는 정보가 있어도 신고를 미루지 마세요."),
            ("전화나 문자를 지워도 될까요?", "**문자와 통화 기록을 우선 남겨 두세요.** 상대방 전화번호, 받은 링크, 대화 내용과 송금 내역을 보관하고 어떤 말을 듣고 송금했는지 시간 순서대로 적어 보세요.\n\n의심스러운 링크를 다시 누르거나 상대방에게 추가로 돈을 보내지는 마세요. 받은 안내가 진짜 기관 연락인지 헷갈리면 해당 기관의 공식 번호로 직접 확인하세요."),
        ],
    },
    "폭행·협박 피해": {
        "source": "https://www.police.go.kr/",
        "questions": [
            ("지금 위협을 받고 있으면 어떻게 해야 하나요?", "**먼저 안전한 곳으로 이동하고 긴급하면 112에 신고하세요.** 다쳤다면 필요한 진료를 받으시고, 직접 상대방을 만나 해결하려고 위험한 장소로 돌아가지 마세요.\n\n신고할 때 현재 위치, 누가 어떤 행동을 했는지, 부상 여부를 있는 그대로 설명하시면 됩니다."),
            ("나중에 상황을 설명하려면 무엇을 남겨야 하나요?", "**사건이 일어난 순서를 짧게 적어 두세요.** 언제·어디서·누가 무엇을 했는지와 현장에서 본 사람이 있는지를 기록하고, 사진·문자·통화 기록·진료 자료 등 이미 가진 자료를 보관하세요.\n\n기억나지 않는 말을 새로 만들어 적지 말고, 직접 겪은 일과 다른 사람에게 들은 이야기를 구분해 주세요."),
        ],
    },
    "교제·가정폭력 피해": {
        "source": "https://www.mogef.go.kr/cc/wcc/cc_wcc_f001.do",
        "questions": [
            ("가까운 사람에게 위협받고 있습니다. 어디에 도움을 요청하나요?", "**지금 위험하다면 안전한 장소로 이동하고 112에 신고하세요.** 상담이나 긴급 보호가 필요하면 **1366**에도 연락할 수 있습니다. 가족이나 연인 사이의 일이라도 혼자 감당할 필요는 없습니다.\n\n안전하게 통화하기 어렵다면 믿을 수 있는 사람에게 대신 도움을 요청하는 방법도 생각해 보세요."),
            ("상황을 어떻게 정리하면 좋을까요?", "**반복된 일을 날짜별로 나누어 적어 보세요.** 폭언·위협·신체 접촉 등 실제 있었던 행동과 당시 장소, 주변에 있었던 사람을 적고, 보유한 대화·사진·진료 자료를 보관하세요.\n\n상대방에게 들킬 위험이 있다면 기록을 남기는 방법보다 본인의 안전을 먼저 생각하고 112 또는 1366에 상황을 설명해 주세요."),
        ],
    },
    "불법촬영·영상 유포 피해": {
        "source": "https://d4u.stop.or.kr/ko/support/content/guide",
        "questions": [
            ("사진이나 영상이 퍼질까 봐 두렵습니다. 어디에 연락하나요?", "**중앙디지털성범죄피해자지원센터에 상담을 요청할 수 있습니다.** 상담 후 상황에 따라 삭제지원과 유포 모니터링 등에 대한 안내를 받을 수 있습니다. **1366**에서도 상담할 수 있고, 긴급한 위협을 받고 있다면 112에 신고하세요.\n\n피해를 확인하려고 영상을 다른 사람에게 보내거나 공개 게시판에 올리지 마세요."),
            ("유포된 게시물을 발견하면 무엇을 남기나요?", "**게시물의 주소(URL), 발견 시각, 게시자 표시를 기록하세요.** 이미 가진 대화나 유포 협박 메시지도 보관해 두시면 상담할 때 설명하기 쉽습니다.\n\n스스로 계속 검색하면서 위험한 자료를 수집해야 하는 것은 아닙니다. 삭제지원이 필요한지 센터와 먼저 상담해 주세요."),
        ],
    },
    "직장 내 괴롭힘": {
        "source": "https://labor.moel.go.kr/minwonApply/minwonFormat.do?searchVal=SN001",
        "questions": [
            ("직장에서 반복적으로 괴롭힘을 당합니다. 어떻게 정리하나요?", "**일어난 일마다 날짜·장소·행동을 적어 보세요.** 누가 어떤 말을 했는지, 업무와 어떤 관련이 있었는지, 주변에 누가 있었는지를 사실 중심으로 남기면 도움이 됩니다.\n\n이미 받은 메시지나 업무 지시, 상담 기록도 보관하세요. 단순히 ‘괴롭힘이다’라고 결론짓기보다 구체적인 행동을 적는 게 상황을 설명하기 쉽습니다."),
            ("어디에 상담하거나 알릴 수 있나요?", "**직장 내 신고 창구를 확인하고, 고용노동부 노동포털의 진정 안내도 살펴보세요.** 누구에게 언제 알렸는지, 어떤 답을 받았는지도 기록해 두시면 좋습니다.\n\n직장 내 괴롭힘에 해당하는지는 실제 상황과 사업장 요건 등에 따라 판단이 달라질 수 있으므로 담당 기관에 구체적인 내용을 설명해 보세요."),
        ],
    },
    "임금·퇴직금 미지급": {
        "source": "https://labor.moel.go.kr/minwonSysInfo/wagesolway.do",
        "questions": [
            ("월급이나 퇴직금을 못 받았습니다. 무엇부터 확인하나요?", "**원래 받기로 한 돈과 실제 받은 돈을 나란히 적어 보세요.** 근무 기간, 월급날, 약속한 금액과 입금된 금액을 정리하고 근로계약서·급여명세서·입금 내역·근무 기록을 보관하세요.\n\n정확한 금액을 아직 모르더라도, 현재 확인되는 내역부터 적어 두시면 상담할 때 도움이 됩니다."),
            ("어디에 신고하거나 상담할 수 있나요?", "**고용노동부 노동포털에서 임금체불 진정을 신청하거나 관할 고용노동관서를 방문해 상담할 수 있습니다.** 임금이 지급되지 않은 날짜와 금액, 사업장 정보와 보유한 자료를 준비해 보세요.\n\n진정 뒤 조사나 출석 안내가 올 수 있으니 담당 기관에서 보낸 연락을 확인해 주세요."),
        ],
    },
    "개인정보 유출·도용": {
        "source": "https://www.privacy.go.kr/front/contents/cntntsView.do?contsNo=45",
        "questions": [
            ("내 정보가 유출됐거나 도용된 것 같습니다. 먼저 뭘 하나요?", "**어디에서 어떤 정보가 노출됐는지 확인되는 범위만 적어 두세요.** 유출 안내 문자나 메일, 내 이름으로 온 낯선 가입·결제 알림, 관련 화면을 보관하고 해당 서비스의 공식 고객센터에 계정 보호 방법을 문의해 보세요.\n\n의심스러운 링크에 개인정보를 다시 입력하지 마세요. 실제 금전 피해가 진행 중이라면 112에도 바로 알리세요."),
            ("개인정보 침해는 어디에 문의하나요?", "**개인정보침해신고센터에 118로 상담하거나 온라인으로 신고할 수 있습니다.** 어떤 서비스에서 무슨 일이 있었는지, 언제 알게 됐는지와 관련 자료를 함께 정리해 두시면 됩니다.\n\n유출 사실이 아직 확인되지 않았다면 ‘유출됐다’고 단정하기보다 발견한 알림과 이상 징후를 그대로 설명해 주세요."),
        ],
    },
}

PASSWORD_RESET_ENABLED = False  # 기본 이메일 템플릿(커스텀 SMTP 미설정)에서는 재설정 링크가 동작하지 않음

UNSUPPORTED_ASSISTANT_MESSAGE = "죄송합니다. 현재 고소음이는 일정 관리·메모 기억 기능을 지원하지 않습니다. 법률 상담이나 사건 정리를 도와드릴게요."

PRIVACY_POLICY_ENABLED = os.environ.get("PRIVACY_POLICY_ENABLED") == "true"
PRIVACY_OPERATOR_NAME = os.environ.get("PRIVACY_OPERATOR_NAME")
PRIVACY_CONTACT_EMAIL = os.environ.get("PRIVACY_CONTACT_EMAIL")
PRIVACY_POLICY_EFFECTIVE_DATE = (os.environ.get("PRIVACY_POLICY_EFFECTIVE_DATE") or "").strip() or None
TERMS_EFFECTIVE_DATE = "2026-09-15"

EVIDENCE_UPLOAD_ENABLED = os.environ.get("EVIDENCE_UPLOAD_ENABLED") == "true"
EVIDENCE_BUCKET = "case-evidence"
EVIDENCE_MAX_FILE_BYTES = 10 * 1024 * 1024
EVIDENCE_MAX_FILES_PER_CASE = 20
EVIDENCE_ALLOWED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif", ".pdf", ".mp3", ".m4a", ".wav"}

NAVER_LOGIN_ENABLED = os.environ.get("NAVER_LOGIN_ENABLED") == "true"
NAVER_LOGIN_START_URL = os.environ.get("NAVER_LOGIN_START_URL")
# 개인정보처리방침에 네이버 수집 항목을 반영한 날. 시행일 환경변수가 이 날짜 이상이고, 한국 날짜로 시행일 당일부터 버튼을 공개한다.
NAVER_PRIVACY_TEXT_REVISED_DATE = datetime.date(2026, 9, 24)
NAVER_ERROR_MESSAGES = {
    "disabled": "네이버 로그인이 현재 준비 중입니다.",
    "cancelled": "네이버 로그인이 취소되었습니다.",
    "invalid_state": "로그인 요청이 만료되었거나 올바르지 않습니다. 네이버 로그인을 처음부터 다시 시도해 주세요.",
    "email_required": "네이버 로그인으로 가입하려면 네이버에서 이메일 제공에 동의해 주세요.",
    "email_exists": "이미 이메일로 가입된 주소입니다. 이메일과 비밀번호로 로그인해 주세요.",
    "server_error": "네이버 로그인 중 일시적인 오류가 발생했습니다. 잠시 후 다시 시도해 주세요.",
}

# Google 로그인은 Supabase OAuth(PKCE)로 처리한다. Client ID/Secret은 Supabase 대시보드에만 둔다.
GOOGLE_LOGIN_ENABLED = os.environ.get("GOOGLE_LOGIN_ENABLED", "true") != "false"
GOOGLE_OAUTH_REDIRECT_URL = (
    os.environ.get("GOOGLE_OAUTH_REDIRECT_URL") or "https://gosoemi-demo-csapxr3rjt2sjg66te9wrt.streamlit.app"
).rstrip("/")
# 개인정보처리방침에 Google 수집 항목을 반영한 날. 시행일 환경변수가 이 날짜 이상이고, 한국 날짜로 시행일 당일부터 버튼을 공개한다.
GOOGLE_PRIVACY_TEXT_REVISED_DATE = datetime.date(2026, 9, 27)
GOOGLE_FLOW_TTL_SECONDS = 600
GOOGLE_FLOW_MAX_PENDING = 2000
GOOGLE_ERROR_MESSAGES = {
    "cancelled": "Google 로그인이 취소되었습니다.",
    "invalid_state": "로그인 요청이 만료되었거나 올바르지 않습니다. Google 로그인을 처음부터 다시 시도해 주세요.",
    "server_error": "Google 로그인 중 일시적인 오류가 발생했습니다. 잠시 후 다시 시도해 주세요.",
}
# 개인정보처리방침 시행일은 한국 날짜 기준으로 판단한다. 서버(Streamlit Cloud)는 UTC로 동작하므로
# date.today()를 쓰면 KST 00:00~09:00 사이에 하루 늦게 활성화된다. 한국은 서머타임이 없어 고정 오프셋을 쓴다.
KST = datetime.timezone(datetime.timedelta(hours=9), "Asia/Seoul")
# 소셜 로그인 사용자의 만 14세 이상 확인 시각을 저장하는 user_metadata 키 (Google 전용, 네이버는 별도 테이블 사용)
SOCIAL_AGE_CONFIRMED_KEY = "age_14_confirmed_at"
SOCIAL_LOGIN_NOTICE = (
    "소셜 계정으로 처음 로그인하시면 만 14세 이상 확인과 이용약관·개인정보 수집·이용 동의를 마친 뒤 "
    "서비스를 이용할 수 있습니다. 제공받는 정보는 개인정보처리방침에서 확인할 수 있습니다."
)


def privacy_config_ready() -> bool:
    return bool(
        PRIVACY_POLICY_ENABLED
        and PRIVACY_OPERATOR_NAME
        and PRIVACY_CONTACT_EMAIL
        and PRIVACY_POLICY_EFFECTIVE_DATE
    )


def terms_config_ready() -> bool:
    return bool(
        PRIVACY_OPERATOR_NAME
        and PRIVACY_CONTACT_EMAIL
        and TERMS_EFFECTIVE_DATE
    )

SERVICE_IDENTITY_INSTRUCTION = (
    "- 이 서비스의 이름은 \"고소음이\"이다.\n"
    "- \"고소음이\"는 소음(noise) 관련 서비스가 아니다. 이름에 \"소음\"이 포함되어 있다고 해서 "
    "소음 측정, 소음 진단, 층간소음 상담 등 소음 관련 서비스로 오해하거나 그렇게 답변하지 마라.\n"
    "- \"고소음이\"는 범죄·피해 상황을 사용자가 정리하고, 고소장 참고용 초안을 작성하도록 돕는 "
    "AI 법률정보 도우미이다.\n"
    "- \"고소음이\"는 변호사의 법률자문을 대체하지 않는다.\n"
    "- 사용자가 서비스의 정체성(무엇을 하는 서비스인지 등)을 물으면 위 내용에 맞게 간결하게 설명하라."
)

DOCX_TITLE = "고소장"
DOCX_SECTION_HEADERS = [
    "고소인", "피고소인", "고소 취지", "범죄 사실", "고소 이유",
    "증거자료", "목격자", "처벌 의사", "작성일",
]


def load_memory():
    if os.path.exists(MEMORY_FILE):
        with open(MEMORY_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return []


def save_memory(memory):
    with open(MEMORY_FILE, "w", encoding="utf-8") as f:
        json.dump(memory, f, ensure_ascii=False, indent=2)


def extract_memory_text(question):
    text = question
    for trigger in MEMORY_TRIGGERS:
        text = re.sub(trigger + r"\S*", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ,.!?？")


def load_schedules():
    if os.path.exists(SCHEDULE_FILE):
        with open(SCHEDULE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return []


def save_schedules(schedules):
    with open(SCHEDULE_FILE, "w", encoding="utf-8") as f:
        json.dump(schedules, f, ensure_ascii=False, indent=2)


def get_supabase_anon_client() -> Client:
    supabase_url = os.environ.get("SUPABASE_URL")
    supabase_key = os.environ.get("SUPABASE_PUBLISHABLE_KEY")
    if not supabase_url or not supabase_key:
        raise RuntimeError("SUPABASE_URL / SUPABASE_PUBLISHABLE_KEY 환경변수가 설정되지 않았습니다.")
    return create_client(supabase_url, supabase_key)


def get_authed_client() -> Client:
    session = st.session_state.get("supabase_session")
    if not session:
        raise RuntimeError("로그인이 필요합니다.")
    client = get_supabase_anon_client()
    client.postgrest.auth(session["access_token"])
    return client


def has_current_consent(user_id: str) -> bool:
    try:
        client = get_authed_client()
        response = (
            client.table("consent_records")
            .select("terms_version")
            .eq("user_id", user_id)
            .eq("terms_version", TERMS_EFFECTIVE_DATE)
            .eq("privacy_version", PRIVACY_POLICY_EFFECTIVE_DATE)
            .limit(1)
            .execute()
        )
        return bool(response.data)
    except Exception:
        return False


def get_app_base_url():
    url = os.environ.get("APP_BASE_URL")
    return url.rstrip("/") if url else None


def log_auth_error(error: Exception) -> None:
    if isinstance(error, AuthApiError):
        print(f"[auth_error] code={error.code} status={error.status}")
    else:
        print(f"[auth_error] unexpected={type(error).__name__}")


def map_auth_error(error: Exception) -> str:
    if isinstance(error, AuthApiError):
        code = getattr(error, "code", None)
        status = getattr(error, "status", None)
        if code == "invalid_credentials":
            return "이메일 또는 비밀번호가 올바르지 않습니다."
        if code == "email_not_confirmed":
            return "이메일 인증이 완료되지 않았습니다. 받은 편지함과 스팸함에서 인증 메일을 확인해 주세요."
        if code in ("user_already_exists", "email_exists"):
            return "이미 가입된 이메일입니다. 로그인하시거나 비밀번호를 재설정해 주세요."
        if code in ("over_email_send_rate_limit", "over_request_rate_limit") or status == 429:
            return "요청이 너무 많습니다. 잠시 후 다시 시도해 주세요."
        if code == "weak_password":
            return "비밀번호가 너무 약합니다. 6자 이상의 다른 비밀번호를 사용해 주세요."
        if code in ("validation_failed", "email_address_invalid"):
            return "이메일 형식이 올바르지 않습니다."
        if code == "otp_expired":
            return "인증 링크가 만료되었거나 이미 사용되었습니다. 인증 메일을 다시 요청해 주세요."
        if code in ("session_not_found", "flow_state_not_found", "flow_state_expired"):
            return "인증 세션이 만료되었습니다. 처음부터 다시 시도해 주세요."
        if code == "same_password":
            return "이전과 동일한 비밀번호로는 변경할 수 없습니다."
        if code == "signup_disabled":
            return "현재 회원가입이 제한되어 있습니다. 잠시 후 다시 시도해 주세요."
    return "일시적인 오류가 발생했습니다. 잠시 후 다시 시도해 주세요."


def set_auth_flash(level, text) -> None:
    st.session_state.auth_flash = {"level": level, "text": text}


def render_auth_flash() -> None:
    flash = st.session_state.get("auth_flash")
    if flash:
        if flash["level"] == "success":
            st.success(flash["text"])
        else:
            st.error(flash["text"])
        st.session_state.auth_flash = None


def session_to_dict(session) -> dict:
    return {
        "access_token": session.access_token,
        "refresh_token": session.refresh_token,
        "user_id": session.user.id,
        "email": session.user.email,
        "expires_at": session.expires_at,
    }


def session_from_auth_response(auth_response) -> dict:
    return session_to_dict(auth_response.session)


# Streamlit Session State is lost on a full browser reload. Store only the
# Supabase refresh token in this browser (never in the URL or a shared cache).
# The component runs in the app origin, and removes the token on logout.
_browser_session_component = st.components.v2.component(
    "gosoemi_browser_session",
    js="""
export default function({data, setStateValue}) {
  const key = 'gosoemi_login_v1';
  try {
    if (data.action === 'clear') {
      localStorage.removeItem(key);
      setStateValue('stored', 'cleared');
    } else if (data.action === 'save') {
      const raw = localStorage.getItem(key);
      let old = null;
      try { old = raw ? JSON.parse(raw) : null; } catch (_) {}
      if (!old || old.token !== data.token) {
        localStorage.setItem(key, JSON.stringify({token: data.token,
          expires: Date.now() + 7 * 24 * 60 * 60 * 1000}));
      }
    } else {
      const raw = localStorage.getItem(key);
      let saved = raw ? JSON.parse(raw) : null;
      if (saved && (typeof saved.token !== 'string' ||
                    !Number.isFinite(saved.expires) || saved.expires < Date.now())) {
        localStorage.removeItem(key);
        saved = null;
      }
      setStateValue('stored', saved ? saved.token : 'empty');
    }
  } catch (_) {
    setStateValue('stored', 'empty');
  }
}
""",
)


def restore_browser_session() -> None:
    if st.session_state.get("browser_login_clear"):
        result = _browser_session_component(
            key="browser_login_clear_component", data={"action": "clear"}
        )
        if result.stored == "cleared":
            st.session_state.browser_login_clear = False
            st.rerun()
        st.info("로그아웃하는 중입니다.")
        st.stop()

    session = st.session_state.get("supabase_session")
    if session:
        _browser_session_component(
            key="browser_login_save_component",
            data={"action": "save", "token": session["refresh_token"]},
        )
        return

    if st.session_state.get("browser_login_checked"):
        return
    result = _browser_session_component(
        key="browser_login_restore_component", data={"action": "restore"}
    )
    token = result.stored
    if token is None:
        st.info("로그인 상태를 확인하는 중입니다.")
        st.stop()
    st.session_state.browser_login_checked = True
    if token not in ("empty", "cleared"):
        try:
            response = get_supabase_anon_client().auth.refresh_session(token)
            if response.session and response.user:
                st.session_state.supabase_session = session_from_auth_response(response)
                st.rerun()
        except Exception as error:
            log_auth_error(error)
        st.session_state.browser_login_clear = True
        st.rerun()


def get_session_authed_client(session: dict) -> Client:
    client = get_supabase_anon_client()
    set_response = client.auth.set_session(session["access_token"], session["refresh_token"])
    if not getattr(set_response, "session", None):
        raise RuntimeError("세션을 설정할 수 없습니다.")
    return client


def clear_user_session_data() -> None:
    st.session_state.case_draft = None
    st.session_state.case_draft_id = None
    st.session_state.last_saved_case_entry = None
    st.session_state.last_saved_fingerprint = None
    st.session_state.last_saved_at = None
    st.session_state.pending_draft_entry = None
    st.session_state.messages = []
    st.session_state.model_history = []
    st.session_state.case_workflow_user_id = None


def perform_logout() -> None:
    session = st.session_state.get("supabase_session")
    if session:
        try:
            get_session_authed_client(session).auth.sign_out()
        except Exception as error:
            log_auth_error(error)
    st.session_state.supabase_session = None
    st.session_state.browser_login_clear = True
    st.session_state.pop("browser_login_clear_component", None)
    st.session_state.pop("google_login_url", None)
    clear_user_session_data()
    st.session_state.auth_view = "login"
    st.rerun()


def ensure_active_session():
    session = st.session_state.get("supabase_session")
    if not session:
        return None
    expires_at = session.get("expires_at")
    if expires_at and time.time() < expires_at - 60:
        return session
    try:
        client = get_supabase_anon_client()
        refreshed = client.auth.refresh_session(session["refresh_token"])
        if not getattr(refreshed, "session", None) or not getattr(refreshed, "user", None):
            raise RuntimeError("empty refresh response")
        new_session = session_from_auth_response(refreshed)
        st.session_state.supabase_session = new_session
        # A refresh token is single-use; write the rotated token to this
        # browser before continuing with the newly refreshed session.
        st.rerun()
    except Exception as error:
        log_auth_error(error)
        st.session_state.supabase_session = None
        st.session_state.browser_login_clear = True
        clear_user_session_data()
        st.session_state.auth_view = "login"
        set_auth_flash("error", "세션이 만료되었습니다. 다시 로그인해 주세요.")
        return None


def handle_auth_query_params() -> None:
    token_hash = st.query_params.get("token_hash")
    type_ = st.query_params.get("type")
    if not token_hash or not type_:
        return

    st.query_params.clear()

    try:
        auth_client = get_supabase_anon_client()
        response = auth_client.auth.verify_otp({"token_hash": token_hash, "type": type_})
        if not getattr(response, "session", None) or not getattr(response, "user", None):
            raise RuntimeError("empty verify_otp response")

        if type_ == "recovery":
            st.session_state.pending_recovery_session = {
                "access_token": response.session.access_token,
                "refresh_token": response.session.refresh_token,
            }
            st.session_state.auth_view = "recovery_set_password"
        else:
            st.session_state.supabase_session = session_from_auth_response(response)
            st.session_state.auth_view = "login"
            set_auth_flash("success", "이메일 인증이 완료되어 자동으로 로그인되었습니다.")
    except Exception as error:
        log_auth_error(error)
        st.session_state.auth_view = "login"
        set_auth_flash("error", map_auth_error(error))

    st.rerun()


def today_kst() -> datetime.date:
    return datetime.datetime.now(KST).date()


def parse_privacy_effective_date():
    try:
        return datetime.date.fromisoformat(PRIVACY_POLICY_EFFECTIVE_DATE or "")
    except ValueError:
        return None


def social_login_gate_checks() -> dict:
    """소셜 로그인 버튼 표시 조건별 통과 여부. 값 자체가 아니라 true/false만 담는다."""
    effective_date = parse_privacy_effective_date()
    return {
        "privacy_enabled": PRIVACY_POLICY_ENABLED,
        "operator_name_present": bool(PRIVACY_OPERATOR_NAME),
        "contact_email_present": bool(PRIVACY_CONTACT_EMAIL),
        "effective_date_valid": effective_date is not None,
        "effective_date_reached_kst": effective_date is not None and today_kst() >= effective_date,
        "naver_policy_date_ok": effective_date is not None and effective_date >= NAVER_PRIVACY_TEXT_REVISED_DATE,
        "naver_enabled": NAVER_LOGIN_ENABLED,
        "naver_start_url_present": bool(NAVER_LOGIN_START_URL),
        "google_policy_date_ok": effective_date is not None and effective_date >= GOOGLE_PRIVACY_TEXT_REVISED_DATE,
        "google_enabled": GOOGLE_LOGIN_ENABLED,
    }


# (조건 키, 실패 사유 코드). 시행일이 없거나 형식이 잘못되면 날짜 관련 사유는
# effective_date_missing / effective_date_invalid 하나로 묶는다.
COMMON_GATE_REASONS = [
    ("privacy_enabled", "privacy_disabled"),
    ("operator_name_present", "operator_name_missing"),
    ("contact_email_present", "contact_email_missing"),
    ("effective_date_valid", "effective_date_invalid"),
    ("effective_date_reached_kst", "effective_date_not_reached_kst"),
]
NAVER_GATE_REASONS = COMMON_GATE_REASONS + [
    ("naver_policy_date_ok", "naver_policy_date_before_revision"),
    ("naver_enabled", "naver_disabled"),
    ("naver_start_url_present", "naver_start_url_missing"),
]
GOOGLE_GATE_REASONS = COMMON_GATE_REASONS + [
    ("google_policy_date_ok", "google_policy_date_before_revision"),
    ("google_enabled", "google_disabled"),
]
DATE_DEPENDENT_GATES = {"effective_date_reached_kst", "naver_policy_date_ok", "google_policy_date_ok"}


def gate_failure_reasons(checks: dict, gates: list) -> list:
    reasons = []
    for key, reason in gates:
        if checks[key]:
            continue
        if key in DATE_DEPENDENT_GATES and not checks["effective_date_valid"]:
            continue
        if reason == "effective_date_invalid" and not PRIVACY_POLICY_EFFECTIVE_DATE:
            reason = "effective_date_missing"
        reasons.append(reason)
    return reasons


def naver_login_available() -> bool:
    return not gate_failure_reasons(social_login_gate_checks(), NAVER_GATE_REASONS)


def google_login_available() -> bool:
    return not gate_failure_reasons(social_login_gate_checks(), GOOGLE_GATE_REASONS)


def format_auth_gate_debug(checks: dict, google_oauth_url_created: bool, google_oauth_attempted: bool) -> str:
    fields = {**checks, "google_oauth_url_created": google_oauth_url_created}
    naver_reasons = gate_failure_reasons(checks, NAVER_GATE_REASONS)
    google_reasons = gate_failure_reasons(checks, GOOGLE_GATE_REASONS)
    if google_oauth_attempted and not google_oauth_url_created:
        google_reasons.append("google_oauth_url_failed")
    parts = [f"{key}={str(value).lower()}" for key, value in fields.items()]
    parts.append(f"naver_visible={str(not naver_reasons).lower()}")
    parts.append(f"naver_hidden_reason={','.join(naver_reasons) or 'none'}")
    parts.append(f"google_visible={str(not google_reasons).lower()}")
    parts.append(f"google_hidden_reason={','.join(google_reasons) or 'none'}")
    return "[auth_gate_debug] " + " ".join(parts)


@st.cache_resource
def get_auth_gate_debug_seen() -> set:
    # 프로세스당 처음 한 번, 그리고 판정 결과가 바뀔 때만 출력해 rerun마다 로그가 쌓이지 않게 한다.
    return set()


def log_auth_gate_debug(google_oauth_url_created: bool, google_oauth_attempted: bool) -> None:
    line = format_auth_gate_debug(social_login_gate_checks(), google_oauth_url_created, google_oauth_attempted)
    seen = get_auth_gate_debug_seen()
    if line in seen:
        return
    seen.add(line)
    print(line, flush=True)


def handle_naver_query_params() -> None:
    naver_token_hash = st.query_params.get("naver_token_hash")
    naver_error = st.query_params.get("naver_error")
    if not naver_token_hash and not naver_error:
        return

    st.query_params.clear()
    st.session_state.auth_view = "login"

    if naver_error:
        set_auth_flash("error", NAVER_ERROR_MESSAGES.get(naver_error, NAVER_ERROR_MESSAGES["server_error"]))
        st.rerun()

    try:
        auth_client = get_supabase_anon_client()
        response = auth_client.auth.verify_otp({"token_hash": naver_token_hash, "type": "magiclink"})
        if not getattr(response, "session", None) or not getattr(response, "user", None):
            raise RuntimeError("empty verify_otp response")
        st.session_state.supabase_session = session_from_auth_response(response)
        set_auth_flash("success", "네이버 계정으로 로그인되었습니다.")
    except Exception as error:
        log_auth_error(error)
        if isinstance(error, AuthApiError) and getattr(error, "code", None) == "otp_expired":
            set_auth_flash("error", NAVER_ERROR_MESSAGES["invalid_state"])
        else:
            set_auth_flash("error", map_auth_error(error))

    st.rerun()


def decode_access_token_claims(session: dict) -> dict:
    # 서명 검증 없이 화면 분기용으로만 사용한다. 실제 권한은 RLS와 RPC가 auth.uid()로 판단한다.
    try:
        payload_segment = session["access_token"].split(".")[1]
        padded = payload_segment + "=" * (-len(payload_segment) % 4)
        return json.loads(base64.urlsafe_b64decode(padded))
    except Exception:
        return {}


def is_naver_session(session: dict) -> bool:
    claims = decode_access_token_claims(session)
    return (claims.get("app_metadata") or {}).get("naver_linked") is True


class GoogleFlowStore:
    """OAuth 시작 시 만든 PKCE code_verifier를 서버 메모리에 잠시 보관한다.

    Google에서 돌아오면 새 Streamlit 세션이 열려 session_state가 비어 있으므로,
    redirect URL에 넣은 일회용 flow ID(해시로만 저장)로 verifier를 찾는다.
    """

    def __init__(self):
        self._lock = threading.Lock()
        self._flows = {}

    @staticmethod
    def _key(flow_id: str) -> str:
        return hashlib.sha256(flow_id.encode("utf-8")).hexdigest()

    def put(self, flow_id: str, code_verifier: str) -> None:
        now = time.time()
        with self._lock:
            for key in [k for k, (_, expires_at) in self._flows.items() if expires_at <= now]:
                del self._flows[key]
            while len(self._flows) >= GOOGLE_FLOW_MAX_PENDING:
                del self._flows[min(self._flows, key=lambda k: self._flows[k][1])]
            self._flows[self._key(flow_id)] = (code_verifier, now + GOOGLE_FLOW_TTL_SECONDS)

    def pop(self, flow_id: str):
        with self._lock:
            entry = self._flows.pop(self._key(flow_id), None)
        if not entry or entry[1] <= time.time():
            return None
        return entry[0]


@st.cache_resource
def get_google_flow_store() -> GoogleFlowStore:
    return GoogleFlowStore()


def get_google_login_url():
    # 같은 브라우저 세션에서는 만료 전까지 같은 URL을 재사용해 rerun마다 flow가 쌓이지 않게 한다.
    cached = st.session_state.get("google_login_url")
    if cached and time.time() < cached["expires_at"] - 60:
        return cached["url"]

    try:
        flow_id = secrets.token_urlsafe(32)
        redirect_to = f"{GOOGLE_OAUTH_REDIRECT_URL}/?{urlencode({'google_flow': flow_id})}"
        auth_client = get_supabase_anon_client()
        response = auth_client.auth.sign_in_with_oauth({
            "provider": "google",
            "options": {
                "redirect_to": redirect_to,
                "query_params": {"prompt": "select_account"},
            },
        })
        storage = auth_client.auth._storage
        code_verifier = storage.get_item(f"{auth_client.auth._storage_key}-code-verifier")
        if not code_verifier or not getattr(response, "url", None):
            raise RuntimeError("missing pkce verifier")
    except Exception as error:
        log_auth_error(error)
        return None

    get_google_flow_store().put(flow_id, code_verifier)
    st.session_state.google_login_url = {
        "url": response.url,
        "expires_at": time.time() + GOOGLE_FLOW_TTL_SECONDS,
    }
    return response.url


def handle_google_query_params() -> None:
    flow_id = st.query_params.get("google_flow")
    if not flow_id:
        return

    auth_code = st.query_params.get("code")
    oauth_error = st.query_params.get("error")
    # 코드가 URL에 남아 새로고침 시 재사용되지 않도록 먼저 지운다.
    st.query_params.clear()
    st.session_state.auth_view = "login"
    code_verifier = get_google_flow_store().pop(flow_id)

    if oauth_error or not auth_code:
        print(f"[auth_error] google_oauth error={oauth_error or 'missing_code'}")
        key = "cancelled" if oauth_error == "access_denied" else "server_error"
        set_auth_flash("error", GOOGLE_ERROR_MESSAGES[key])
        st.rerun()

    if not code_verifier:
        set_auth_flash("error", GOOGLE_ERROR_MESSAGES["invalid_state"])
        st.rerun()

    try:
        auth_client = get_supabase_anon_client()
        response = auth_client.auth.exchange_code_for_session(
            {"auth_code": auth_code, "code_verifier": code_verifier}
        )
        if not getattr(response, "session", None) or not getattr(response, "user", None):
            raise RuntimeError("empty exchange_code_for_session response")
        st.session_state.supabase_session = session_from_auth_response(response)
        st.session_state.pop("google_login_url", None)
        set_auth_flash("success", "Google 계정으로 로그인되었습니다.")
    except Exception as error:
        log_auth_error(error)
        set_auth_flash("error", map_auth_error(error))

    st.rerun()


def is_google_primary_session(session: dict) -> bool:
    # 이메일로 가입한 뒤 같은 이메일의 Google 계정이 연결된 경우는 provider가 "email"로 남아 제외된다.
    claims = decode_access_token_claims(session)
    return (claims.get("app_metadata") or {}).get("provider") == "google"


def social_age_confirmed(session: dict) -> bool:
    claims = decode_access_token_claims(session)
    return bool((claims.get("user_metadata") or {}).get(SOCIAL_AGE_CONFIRMED_KEY))


def confirm_google_age() -> None:
    client = get_session_authed_client(st.session_state.supabase_session)
    client.auth.update_user({
        "data": {SOCIAL_AGE_CONFIRMED_KEY: datetime.datetime.now(datetime.timezone.utc).isoformat()}
    })
    # 새 access token에 갱신된 user_metadata가 담기도록 세션을 다시 발급받는다.
    refreshed = client.auth.refresh_session()
    if not getattr(refreshed, "session", None) or not getattr(refreshed, "user", None):
        raise RuntimeError("empty refresh response")
    st.session_state.supabase_session = session_from_auth_response(refreshed)


DUPLICATE_SUBMIT_WINDOW_SECONDS = 10


def compute_case_fingerprint(user_id: str, field_values: dict) -> str:
    normalized = {key: (field_values.get(key) or "").strip() for _, key in CASE_FIELDS}
    payload = json.dumps({"user_id": user_id, "fields": normalized}, ensure_ascii=True, sort_keys=True)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def save_case(case_entry):
    client = get_authed_client()
    client.table(CASES_TABLE).insert(case_entry).execute()


def load_cases_safe():
    try:
        session = st.session_state.get("supabase_session")
        if not session or not session.get("user_id"):
            raise RuntimeError("로그인이 필요합니다.")
        client = get_authed_client()
        response = (
            client.table(CASES_TABLE)
            .select("*")
            .eq("user_id", session["user_id"])
            .order("created_at", desc=True)
            .execute()
        )
    except Exception:
        return [], "사건 데이터를 불러오는 중 문제가 발생했습니다. 잠시 후 다시 시도해 주세요."
    return response.data or [], None


def evidence_folder(user_id: str, case_id: str) -> str:
    return f"{user_id}/{case_id}"


def safe_evidence_filename(filename: str) -> str:
    basename = os.path.basename((filename or "").replace("\\", "/"))
    stem, extension = os.path.splitext(basename)
    extension = extension.lower() if extension.lower() in EVIDENCE_ALLOWED_EXTENSIONS else ""
    # Supabase Storage 객체 키는 ASCII 일부 문자만 허용하므로 한글 등은 제거한다.
    stem = unicodedata.normalize("NFKD", stem).encode("ascii", "ignore").decode("ascii")
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._-")
    return f"{(stem or 'evidence')[:100]}{extension}"


def evidence_item_value(item, key, default=None):
    if isinstance(item, dict):
        return item.get(key, default)
    return getattr(item, key, default)


def evidence_display_name(object_name: str) -> str:
    if "__" in object_name:
        return object_name.split("__", 1)[1]
    if re.match(r"^[0-9a-f]{32}_", object_name):
        return object_name[33:]
    return object_name


def get_storage_authed_client() -> Client:
    session = ensure_active_session()
    if not session or not session.get("user_id"):
        raise RuntimeError("로그인이 필요합니다.")
    client = get_session_authed_client(session)
    auth_session = client.auth.get_session()
    if not auth_session or not auth_session.user or auth_session.user.id != session["user_id"]:
        raise RuntimeError("세션 사용자 정보가 일치하지 않습니다.")
    if auth_session.access_token != session["access_token"]:
        st.session_state.supabase_session = session_to_dict(auth_session)
    return client


def safe_error_message(error: Exception) -> str:
    message = str(getattr(error, "message", None) or error)
    message = re.sub(r"eyJ[\w-]+\.[\w-]+\.[\w-]+", "[token]", message)
    message = re.sub(r"[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}", "[id]", message)
    message = re.sub(r"[0-9a-fA-F]{32}", "[id]", message)
    message = re.sub(r"[\w.+-]+@[\w-]+\.[\w.-]+", "[email]", message)
    message = re.sub(r"\[id\]/\S+", "[path]", message)
    return " ".join(message.split())[:200]


def log_evidence_upload_error(stage: str, error: Exception) -> None:
    status = getattr(error, "status", None) or getattr(error, "status_code", None)
    code = getattr(error, "code", None)
    print(
        f"[evidence_upload_error] stage={stage} type={type(error).__name__} "
        f"status={status} code={code} message={safe_error_message(error)}",
        flush=True,
    )


def map_evidence_upload_error(error: Exception) -> str:
    status = str(getattr(error, "status", None) or getattr(error, "status_code", None) or "")
    message = str(getattr(error, "message", None) or error).lower()
    if status in ("401", "403") or "row-level security" in message or "unauthorized" in message or "jwt" in message:
        return "업로드 권한이 확인되지 않았습니다. 로그아웃 후 다시 로그인해 주세요."
    if status == "413" or "maximum allowed size" in message or "too large" in message:
        return "파일 크기가 10MB를 초과해 올릴 수 없습니다."
    if status == "409" or "already exists" in message:
        return "같은 파일이 이미 있습니다. 잠시 후 다시 시도해 주세요."
    if "mime" in message or "invalid_mime_type" in message:
        return "지원하지 않는 파일 형식입니다."
    if status == "404" or "bucket not found" in message:
        return "증거자료 저장소를 찾을 수 없습니다. 관리자에게 문의해 주세요."
    if "로그인" in str(error) or "세션" in str(error):
        return "로그인이 만료되었습니다. 다시 로그인한 뒤 시도해 주세요."
    return "업로드에 실패했습니다. 잠시 후 다시 시도해 주세요."


def list_case_evidence(case_id: str):
    session = st.session_state.get("supabase_session")
    if not session or not session.get("user_id"):
        raise RuntimeError("로그인이 필요합니다.")
    folder = evidence_folder(session["user_id"], case_id)
    client = get_storage_authed_client()
    items = client.storage.from_(EVIDENCE_BUCKET).list(
        folder,
        {"limit": EVIDENCE_MAX_FILES_PER_CASE + 1, "sortBy": {"column": "created_at", "order": "desc"}},
    )
    files = []
    for item in items or []:
        name = evidence_item_value(item, "name")
        item_id = evidence_item_value(item, "id")
        if name and item_id:
            metadata = evidence_item_value(item, "metadata", {}) or {}
            files.append({
                "name": name,
                "display_name": evidence_display_name(name),
                "size": evidence_item_value(metadata, "size", 0) or 0,
                "path": f"{folder}/{name}",
            })
    return files


def validate_evidence_upload(uploaded_file) -> None:
    extension = os.path.splitext(uploaded_file.name or "")[1].lower()
    if extension not in EVIDENCE_ALLOWED_EXTENSIONS:
        raise ValueError("지원하지 않는 파일 형식입니다.")
    if uploaded_file.size > EVIDENCE_MAX_FILE_BYTES:
        raise ValueError("파일 한 개의 크기는 10MB 이하여야 합니다.")
    if uploaded_file.size <= 0:
        raise ValueError("빈 파일은 올릴 수 없습니다.")


def upload_case_evidence(case_id: str, uploaded_file) -> None:
    validate_evidence_upload(uploaded_file)
    existing = list_case_evidence(case_id)
    if len(existing) >= EVIDENCE_MAX_FILES_PER_CASE:
        raise ValueError("사건 하나에는 증거자료를 최대 20개까지 첨부할 수 있습니다.")
    client = get_storage_authed_client()
    auth_user = client.auth.get_session().user
    safe_name = safe_evidence_filename(uploaded_file.name)
    object_path = f"{evidence_folder(auth_user.id, case_id)}/{uuid.uuid4().hex}_{safe_name}"
    client.storage.from_(EVIDENCE_BUCKET).upload(
        object_path,
        uploaded_file.getvalue(),
        {"content-type": uploaded_file.type or "application/octet-stream", "cache-control": "3600", "upsert": "false"},
    )


def download_case_evidence(object_path: str):
    session = st.session_state.get("supabase_session")
    if not session or not session.get("user_id") or not object_path.startswith(f"{session['user_id']}/"):
        raise RuntimeError("접근 권한이 없습니다.")
    return get_storage_authed_client().storage.from_(EVIDENCE_BUCKET).download(object_path)


def delete_case_evidence_file(object_path: str) -> None:
    session = st.session_state.get("supabase_session")
    if not session or not session.get("user_id") or not object_path.startswith(f"{session['user_id']}/"):
        raise RuntimeError("접근 권한이 없습니다.")
    get_storage_authed_client().storage.from_(EVIDENCE_BUCKET).remove([object_path])


def delete_all_case_evidence(case_id: str) -> None:
    if not EVIDENCE_UPLOAD_ENABLED:
        return
    files = list_case_evidence(case_id)
    if files:
        get_storage_authed_client().storage.from_(EVIDENCE_BUCKET).remove([item["path"] for item in files])


def delete_case(case_id):
    session = st.session_state.get("supabase_session")
    if not session or not session.get("user_id"):
        raise RuntimeError("로그인이 필요합니다.")
    delete_all_case_evidence(case_id)
    client = get_authed_client()
    client.table(CASES_TABLE).delete().eq("case_id", case_id).eq("user_id", session["user_id"]).execute()


def build_draft_prompt(case, today_str):
    fields_text = "\n".join(
        f"- {label}: {case.get(key) or '(입력 안 됨)'}"
        for label, key in CASE_FIELDS
    )

    return (
        "아래는 사용자가 실제로 입력한 사건 정보입니다. 이 정보만 근거로 고소장 초안을 작성하세요.\n"
        "규칙:\n"
        "1. 사용자가 입력한 객관적인 사실과 사용자의 주장만 작성하세요. 사용자가 제공하지 않은 내용을 추측하거나 새로 만들어내지 마세요.\n"
        "2. 확인되지 않은 범죄 고의, 기망 의도, 불법영득 의사, 범죄 성립 여부를 확정적으로 표현하지 마세요.\n"
        "3. 필요한 경우 \"고소인은 ~라고 판단합니다\", \"~한 것으로 의심됩니다\", \"수사를 통해 확인해 주시기 바랍니다\"와 같이 고소인의 주장·의심임을 분명히 하는 표현을 사용하세요.\n"
        "4. 사용자가 직접 말하지 않은 법조문, 범행 동기, 고의, 추가 피해 내용을 만들어내지 마세요.\n"
        "5. 특정 항목에 필요한 정보가 제공되지 않았거나 불충분하면 반드시 '확인 필요'라고만 표시하세요.\n"
        "6. \"고소 취지\"와 \"범죄 사실\" 항목은 법원의 확정판결처럼 단정하지 말고, 고소인의 주장 형식으로 작성하세요.\n"
        "7. 다음 10개 항목을 정확히 이 순서로, 각 항목을 소제목으로 구분해 작성하세요:\n"
        "   문서 제목 / 고소인 / 피고소인 / 고소 취지 / 범죄 사실 / 고소 이유 / 증거자료 / 목격자 / 처벌 의사 / 작성일\n"
        f"8. 작성일은 반드시 {today_str}로 표기하세요.\n\n"
        f"[입력된 사건 정보]\n{fields_text}"
    )


def generate_case_draft(case):
    kst = datetime.timezone(datetime.timedelta(hours=9))
    today_str = datetime.datetime.now(kst).date().isoformat()

    prompt = build_draft_prompt(case, today_str)
    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
    response = client.models.generate_content(
        model="gemini-3.5-flash-lite",
        contents=[{"role": "user", "parts": [{"text": prompt}]}],
    )

    return f"{response.text}\n\n{DRAFT_DISCLAIMER}"


def build_case_draft_docx(draft_text):
    document = Document()

    section = document.sections[0]
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    normal_style = document.styles["Normal"]
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.line_spacing = 1.5

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(DOCX_TITLE)
    title_run.bold = True
    title_run.font.size = Pt(20)

    document.add_paragraph()

    skip_next_line = False
    for raw_line in draft_text.splitlines():
        line = raw_line.strip()

        if not line:
            continue
        if line == "문서 제목":
            skip_next_line = True
            continue
        if skip_next_line:
            skip_next_line = False
            continue

        if line in DOCX_SECTION_HEADERS:
            heading_paragraph = document.add_paragraph()
            heading_run = heading_paragraph.add_run(line)
            heading_run.bold = True
            heading_run.font.size = Pt(13)
        else:
            document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def get_google_calendar_schedule(question):
    target_date = resolve_date(question)

    if target_date is None:
        target_date = datetime.date.today()

    kst = datetime.timezone(datetime.timedelta(hours=9))

    start_time = datetime.datetime.combine(
        target_date,
        datetime.time.min,
        tzinfo=kst
    )

    end_time = start_time + datetime.timedelta(days=1)

    service = get_calendar_service()

    result = service.events().list(
        calendarId="primary",
        timeMin=start_time.isoformat(),
        timeMax=end_time.isoformat(),
        singleEvents=True,
        orderBy="startTime"
    ).execute()

    events = result.get("items", [])

    if not events:
        return f"{target_date} 일정이 없습니다."

    schedule_text = []

    for event in events:
        start = event["start"].get("dateTime", event["start"].get("date"))
        title = event.get("summary", "제목 없음")
        schedule_text.append(f"- {start} | {title}")

    return "\n".join(schedule_text)
def has_date_cue(question):
    return (
        any(word in question for word in DATE_WORDS)
        or re.search(r"\d{1,2}월\s*\d{1,2}일", question)
        or re.search(r"\d{1,2}시", question)
    )


def resolve_date(question, today=None):
    today = today or datetime.date.today()
    if "모레" in question:
        return today + datetime.timedelta(days=2)
    if "내일" in question:
        return today + datetime.timedelta(days=1)
    if "오늘" in question:
        return today
    m = re.search(r"(\d{1,2})월\s*(\d{1,2})일", question)
    if m:
        month, day = int(m.group(1)), int(m.group(2))
        candidate = datetime.date(today.year, month, day)
        if candidate < today:
            candidate = datetime.date(today.year + 1, month, day)
        return candidate
    return None


def resolve_date_range(question, today=None):
    today = today or datetime.date.today()
    if "다음주" in question or "다음 주" in question:
        this_monday = today - datetime.timedelta(days=today.weekday())
        next_monday = this_monday + datetime.timedelta(days=7)
        return next_monday, next_monday + datetime.timedelta(days=6)
    if "이번주" in question or "이번 주" in question:
        this_monday = today - datetime.timedelta(days=today.weekday())
        return this_monday, this_monday + datetime.timedelta(days=6)
    single_date = resolve_date(question, today)
    if single_date:
        return single_date, single_date
    return None, None


def resolve_time(question):
    m = re.search(r"(오전|오후)?\s*(\d{1,2})\s*시(?:\s*(\d{1,2})\s*분)?", question)    
    if not m:
        return None

    period = m.group(1)
    hour = int(m.group(2))
    minute = int(m.group(3)) if m.group(3) else 0

    if period == "오후" and hour != 12:
        hour += 12
    if period == "오전" and hour == 12:
        hour = 0

    return f"{hour:02d}:{minute:02d}"    


def strip_common_words(text):
    for word in DATE_WORDS:
        text = text.replace(word, "")
    text = re.sub(r"\d{1,2}월\s*\d{1,2}일", "", text)
    text = re.sub(r"(오전|오후)?\s*\d{1,2}시(\s*\d{1,2}분)?", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ,.!?에는를을이가")


def extract_task_text(question):
    text = question
    for trigger in MEMORY_TRIGGERS:
        text = re.sub(trigger + r"\S*", "", text)
    return strip_common_words(text)


def extract_cancel_keyword(question):
    text = question
    for trigger in CANCEL_TRIGGERS:
        text = re.sub(trigger + r"\S*", "", text)
    text = text.replace("일정", "")
    return strip_common_words(text)


def is_cancel_request(question):
    return any(trigger in question for trigger in CANCEL_TRIGGERS)


def is_schedule_add_request(question):
    return any(word in question for word in ("추가", "등록", "저장", "잡아", "만들어", "넣어")) and ("일정" in question or has_date_cue(question))
def is_schedule_query(question):
    return "일정" in question and not is_cancel_request(question)


def format_schedule(item):
    time_part = f" {item['time']}" if item["time"] else ""
    return f"{item['date']}{time_part} {item['task']}"


def handle_schedule_add(question, schedules):
    date = resolve_date(question)
    time = resolve_time(question)
    task = extract_task_text(question)
    print(f"[확인] date={date}, time={time}, task={task!r}")
    if not date or not task:
        print("[일정 저장 실패] 날짜나 내용을 정확히 알 수 없습니다.")
        return "사용자가 일정을 저장하려 했지만 날짜 또는 내용을 인식하지 못했습니다."

    entry = {
        "date": date.isoformat(),
        "time": time,
        "task": task
    }

    schedules.append(entry)
    save_schedules(schedules)

    try:
        service = get_calendar_service()

        if time:
            hour, minute = map(int, time.split(":")[:2])

            start_dt = datetime.datetime(
                date.year, date.month, date.day,
                hour, minute,
                tzinfo=datetime.timezone(datetime.timedelta(hours=9))
            )

            end_dt = start_dt + datetime.timedelta(hours=1)

            event = {
                "summary": task,
                "start": {
                    "dateTime": start_dt.isoformat(),
                    "timeZone": "Asia/Seoul"
                },
                "end": {
                    "dateTime": end_dt.isoformat(),
                    "timeZone": "Asia/Seoul"
                }
            }

        else:
            event = {
                "summary": task,
                "start": {
                    "date": date.isoformat()
                },
                "end": {
                    "date": (date + datetime.timedelta(days=1)).isoformat()
                }
            }

        service.events().insert(
            calendarId="primary",
            body=event
        ).execute()

        return f"일정을 저장했고 구글 캘린더에도 등록했습니다: {format_schedule(entry)}"

    except Exception as e:
        print(f"[구글 캘린더 저장 실패] {e}")
        return f"일정은 저장했지만 구글 캘린더 등록은 실패했습니다: {format_schedule(entry)}"

def handle_schedule_cancel(question, schedules):
    date = resolve_date(question)
    keyword = extract_cancel_keyword(question)

    if not keyword:
        return "어떤 일정을 취소할지 일정 이름을 말씀해 주세요."

    local_matches = [
        s for s in schedules
        if (date is None or s["date"] == date.isoformat())
        and keyword in s["task"]
    ]

    for item in local_matches:
        schedules.remove(item)

    if local_matches:
        save_schedules(schedules)

    kst = datetime.timezone(datetime.timedelta(hours=9))
    target_date = date or datetime.datetime.now(kst).date()

    start_time = datetime.datetime.combine(
        target_date,
        datetime.time.min,
        tzinfo=kst
    )

    if date is None:
        end_time = start_time + datetime.timedelta(days=365)
    else:
        end_time = start_time + datetime.timedelta(days=1)

    google_matches = []

    try:
        service = get_calendar_service()
        result = service.events().list(
            calendarId="primary",
            timeMin=start_time.isoformat(),
            timeMax=end_time.isoformat(),
            singleEvents=True,
            orderBy="startTime"
        ).execute()

        for event in result.get("items", []):
            title = event.get("summary", "")
            if keyword in title:
                service.events().delete(
                    calendarId="primary",
                    eventId=event["id"]
                ).execute()
                google_matches.append(title)

    except Exception as e:
        print(f"[구글 캘린더 삭제 실패] {e}")

    removed_items = [format_schedule(item) for item in local_matches]
    removed_items.extend(google_matches)

    if not removed_items:
        return "조건에 맞는 일정을 찾지 못했습니다. 어떤 일정인지 다시 말씀해 주세요."

    removed_text = ", ".join(dict.fromkeys(removed_items))
    print(f"[일정 삭제됨] {removed_text}")
    return f"요청하신 일정을 삭제했습니다: {removed_text}"


def handle_schedule_query(question, schedules):
    start, end = resolve_date_range(question)
    if start and end:
        matches = [s for s in schedules if start.isoformat() <= s["date"] <= end.isoformat()]
    else:
        today_str = datetime.date.today().isoformat()
        matches = [s for s in schedules if s["date"] >= today_str]
    matches.sort(key=lambda s: (s["date"], s["time"] or ""))
    if not matches:
        print("[조회된 일정 없음]")
        return "요청한 기간에 해당하는 저장된 일정이 없습니다. 없다고 자연스럽게 답변하세요."
    listing = "\n".join(f"- {format_schedule(m)}" for m in matches)
    print(f"[조회된 일정]\n{listing}")
    return f"다음은 요청한 기간에 해당하는 저장된 일정입니다. 이 내용만 근거로 답변하세요:\n{listing}"


def build_system_instruction(memory, schedule_notice):
    parts = [SERVICE_IDENTITY_INSTRUCTION]
    if memory:
        lines = "\n".join(f"- {item}" for item in memory)
        parts.append(
            "다음은 사용자가 이전에 기억해달라고 요청한 장기 기억입니다. "
            "관련된 질문에 답할 때 반드시 참고하세요:\n" + lines
        )
    if schedule_notice:
        parts.append(schedule_notice)
    return "\n\n".join(parts)


def process_question(question, history):
    question = unicodedata.normalize("NFC", question)

    if is_cancel_request(question) and ("일정" in question or has_date_cue(question)):
        return UNSUPPORTED_ASSISTANT_MESSAGE

    if is_schedule_add_request(question):
        return UNSUPPORTED_ASSISTANT_MESSAGE

    if (
        ("캘린더" in question)
        or ("구글 일정" in question)
        or (has_date_cue(question) and ("일정" in question or "스케줄" in question))
    ):
        return UNSUPPORTED_ASSISTANT_MESSAGE

    if any(trigger in question for trigger in MEMORY_TRIGGERS):
        return UNSUPPORTED_ASSISTANT_MESSAGE

    if is_schedule_query(question):
        return UNSUPPORTED_ASSISTANT_MESSAGE

    history.append({"role": "user", "parts": [{"text": question}]})
    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
    response = client.models.generate_content(
        model="gemini-3.5-flash-lite",
        contents=history[-16:],
        config=types.GenerateContentConfig(
            system_instruction=build_system_instruction([], None)
        ),
    )

    history.append({"role": "model", "parts": [{"text": response.text}]})
    if len(history) > 16:
        history[:] = history[-16:]
    return response.text


BRAND_CSS = """
<style>
:root {
    --dgo-navy: #0B1F3A;
    --dgo-navy-light: #13294B;
    --dgo-blue: #2F80ED;
    --dgo-blue-light: #EAF2FE;
    --dgo-muted: #5B6B82;
}

.dgo-hero {
    text-align: center;
    padding: 0.5rem 0 1.25rem 0;
}

.dgo-hero-title {
    font-size: 2.1rem;
    font-weight: 800;
    color: var(--dgo-navy);
    letter-spacing: -0.02em;
}

.dgo-hero-subtitle {
    margin-top: 0.4rem;
    font-size: 1.02rem;
    color: var(--dgo-muted);
}

.dgo-card {
    background-color: var(--dgo-blue-light);
    border: 1px solid #D3E4FB;
    border-radius: 14px;
    padding: 1.1rem 0.8rem;
    text-align: center;
    font-size: 1.4rem;
    color: var(--dgo-navy);
    margin-bottom: 0.75rem;
}

.dgo-card b {
    display: block;
    margin-top: 0.35rem;
    font-size: 0.95rem;
    font-weight: 700;
}

.dgo-footer {
    margin-top: 2rem;
    padding-top: 1rem;
    border-top: 1px solid #E3E8EF;
    text-align: center;
    font-size: 0.82rem;
    color: var(--dgo-muted);
}

[data-testid="stButton"] button,
[data-testid="stFormSubmitButton"] button,
[data-testid="stDownloadButton"] button {
    background-color: var(--dgo-blue);
    color: #FFFFFF;
    border: none;
    border-radius: 8px;
    font-weight: 600;
}

[data-testid="stButton"] button:hover,
[data-testid="stFormSubmitButton"] button:hover,
[data-testid="stDownloadButton"] button:hover {
    background-color: var(--dgo-navy-light);
    color: #FFFFFF;
}

@media (max-width: 640px) {
    .dgo-hero-title {
        font-size: 1.6rem;
    }
    .dgo-hero-subtitle {
        font-size: 0.92rem;
    }
    .dgo-card {
        font-size: 1.2rem;
    }
}

@media (max-width: 768px) {
    html, body {
        -webkit-text-size-adjust: 100%;
        text-size-adjust: 100%;
    }

    .dgo-card {
        padding: 0.6rem 0.7rem;
        margin-bottom: 0.4rem;
    }

    [data-testid="stMarkdownContainer"] p,
    [data-testid="stMarkdownContainer"] li {
        font-size: 0.9rem;
        line-height: 1.5;
    }

    [data-testid="stTextArea"] textarea {
        height: 120px !important;
        min-height: 120px !important;
        resize: vertical;
    }

    [data-testid="stTextInput"] input,
    [data-testid="stButton"] button,
    [data-testid="stFormSubmitButton"] button,
    [data-testid="stDownloadButton"] button {
        min-height: 44px;
    }
}

.dgo-google-btn,
.dgo-google-btn:link,
.dgo-google-btn:visited {
    display: inline-flex;
    align-items: center;
    justify-content: center;
    gap: 10px;
    min-height: 40px;
    padding: 0 16px;
    margin: 0.25rem 0 0.5rem 0;
    background-color: #FFFFFF;
    border: 1px solid #747775;
    border-radius: 8px;
    color: #1F1F1F !important;
    font-family: "Roboto", Arial, sans-serif;
    font-size: 14px;
    font-weight: 500;
    text-decoration: none !important;
}

.dgo-google-btn:hover {
    background-color: #F8F9FA;
    box-shadow: 0 1px 2px rgba(60, 64, 67, 0.3);
}

.dgo-google-btn:focus-visible {
    outline: 2px solid var(--dgo-blue);
    outline-offset: 2px;
}

@media (max-width: 768px) {
    .dgo-google-btn {
        min-height: 44px;
    }
}

#MainMenu {visibility: hidden;}
footer {visibility: hidden;}
[data-testid="stToolbar"] {visibility: hidden;}
[data-testid="stAppDeployButton"] {display: none;}
</style>
"""


# 법률 검토 전 초안 — 실제 배포 전 검토 필요
def build_privacy_policy_text():
    return f"""
### 개인정보처리방침

{PRIVACY_OPERATOR_NAME}(이하 "운영자")는 개인이 운영하는 서비스(사업자등록 전)로서 "고소음이"(이하 "서비스")를 제공하며,
이용자의 개인정보를 아래와 같이 처리합니다.

**1. 수집하는 개인정보 항목**
- 회원가입 시: 이메일 주소, 인증을 위한 비밀번호
- 네이버 로그인 이용 시: 네이버 고유 식별자, 네이버 계정 이메일 (이용자 동의에 따라 네이버로부터 제공받음)
  - 네이버 로그인 과정에서 발급되는 네이버 접근 토큰은 이용자 확인에만 사용하며, 서비스의 사건 데이터베이스에는 저장하지 않습니다. 로그인 직후 네이버 서비스 연결을 해제하지 않습니다.
- Google 로그인 이용 시: Google 고유 식별자, Google 계정 이메일, 이름, 프로필 사진 (이용자 동의에 따라 Google로부터 제공받음)
- 증거자료 첨부 이용 시: 이용자가 직접 선택한 사진, PDF 문서, 녹음파일 및 파일명
- 사건 작성 시: 사건 유형, 신고인 정보, 상대방 정보, 사건 발생 일시, 사건 발생 장소, 사건 경위, 피해 내용, 증거자료, 목격자 유무, 처벌 의사 (총 10개 항목)
  - 이 중 "상대방 정보"는 이용자 본인이 아닌 제3자에 관한 정보가 포함될 수 있습니다.

**2. 개인정보의 수집·이용 목적**
- 회원가입, 로그인, 계정 관리
- 네이버 고유 식별자·네이버 계정 이메일: 네이버 로그인을 통한 회원가입·로그인·계정관리
- Google 고유 식별자·Google 계정 이메일·이름·프로필 사진: Google 로그인을 통한 회원가입·로그인·계정관리
- 이용자가 입력한 사건 정보와 첨부한 증거자료의 비공개 저장·조회·다운로드·삭제
- AI를 이용한 상담 응답 제공 및 고소장 초안 생성

**3. 개인정보의 보유·이용 기간**
- 회원 정보(이메일, 비밀번호, 네이버 고유 식별자, 네이버 계정 이메일, Google 고유 식별자, Google 계정 이메일·이름·프로필 사진): 회원탈퇴 시까지 보유하며, 탈퇴 확인 후 삭제합니다.
- 사건 정보 및 첨부한 증거자료: 이용자가 직접 삭제하거나 회원탈퇴를 요청할 때까지 보유합니다.

**4. 개인정보 처리 위탁 및 정보 전송**
- Supabase — 회원 인증, 사건 정보 및 첨부한 증거자료의 비공개 저장. 데이터베이스와 인증 정보는 AWS 서울 리전에 저장되어 주로 처리됩니다.
  다만 로그·모니터링·이메일 발송 등 부가 처리 과정에서 미국·싱가포르 소재 하위처리자를 거칠 수 있습니다.
  비밀번호는 Supabase Auth에서 암호화된 인증정보로 처리되며 운영자는 원문을 열람할 수 없습니다.
- Google Gemini API(유료 API) — 고소장 초안 생성 및 AI 상담 응답 생성을 위해 입력하신 정보가 처리 목적으로 전송됩니다.
  유료 API 이용 시 입력·응답 내용은 제품 개선(모델 학습 포함)에 사용되지 않으며, 오남용 방지 등을 위한 로그가 기본 최대 55일간 보관될 수 있습니다.

**5. 개인정보의 국외 이전**
- Supabase의 로그·모니터링·이메일 발송 등 부가 처리는 미국·싱가포르 소재 하위처리자를 거칠 수 있습니다.
- Google Gemini API로 전송되는 정보는 Google 또는 그 대리인이 시설을 운영하는 국가에서 일시 저장·캐시될 수 있습니다. 구체적인 처리 국가와 계약 법인은 현재 확인되지 않아 확인되는 대로 갱신하겠습니다.

**6. 이용자의 권리**
- 저장된 사건의 열람·삭제: 서비스 내 "내 사건" 화면에서 직접 삭제할 수 있습니다.
- 회원탈퇴 및 전체 정보 삭제: 아래 문의처로 이메일을 보내 요청하시면, 확인 후 계정과 저장된 모든 사건 정보를 삭제합니다. 현재는 이메일 요청을 받아 순차적으로 처리하고 있습니다.

**7. 이용 연령 제한**
본 서비스는 만 14세 이상만 이용할 수 있습니다.

**8. 안전한 정보 입력 안내**
상대방 등 제3자의 정보는 이용자가 적법하게 제공할 수 있는 범위에서 필요한 최소한만 입력해 주세요.
주민등록번호, 계좌 비밀번호, 카드번호, 인증번호(OTP), 불필요한 주소·연락처 등은 입력하지 마세요.

**9. 문의처**
개인정보 관련 문의, 열람·정정·삭제 요청은 아래로 연락해 주세요.
- 운영자: {PRIVACY_OPERATOR_NAME}
- 이메일: {PRIVACY_CONTACT_EMAIL}

**10. 시행일**
이 개인정보처리방침은 {PRIVACY_POLICY_EFFECTIVE_DATE}부터 적용됩니다.
"""


@st.dialog("개인정보처리방침", width="large")
def render_privacy_policy_dialog():
    if not privacy_config_ready():
        st.info("개인정보처리방침을 준비 중입니다. 잠시 후 다시 확인해 주세요.")
    else:
        st.markdown(build_privacy_policy_text())


def build_terms_of_service_text():
    return f"""
### 이용약관

이 약관은 {PRIVACY_OPERATOR_NAME}(이하 "운영자")가 개인 자격으로 제공하는 "고소음이"(이하 "서비스")의 이용에 관한 조건을 정합니다.
이 약관은 회원가입 과정에서 이용약관 필수 동의 항목에 체크하고 가입을 완료한 이용자에게 적용됩니다.

**제1조 (목적)**
이 약관은 운영자가 제공하는 서비스의 이용조건 및 절차, 운영자와 이용자의 권리·의무 및 책임사항을 규정함을 목적으로 합니다.

**제2조 (서비스의 내용과 법률자문 비대체)**
1. 서비스는 이용자가 입력한 피해 상황을 정리하고, 고소장 등 문서의 참고용 초안을 작성하도록 돕는 법률정보 제공 및 문서작성 보조 서비스입니다.
2. 서비스는 변호사가 아니며, 서비스가 제공하는 정보와 문서는 변호사의 법률자문이나 소송대리를 대체하지 않습니다. 서비스 이용만으로 특정 법적 결과(고소 성립, 수사 개시, 처벌 등)가 보장되지 않습니다.
3. 구체적인 법적 판단이 필요한 경우 이용자는 반드시 변호사 등 관련 분야 전문가와 별도로 상담해야 합니다.

**제3조 (AI 생성 결과의 한계와 전문가 검토 안내)**
1. 서비스가 AI를 통해 생성하는 상담 응답 및 고소장 초안(이하 "AI 생성 결과")은 이용자가 입력한 정보를 바탕으로 자동 생성된 참고 자료이며, 그 정확성·완전성·최신 법령 반영 여부를 보장하지 않습니다.
2. 이용자는 AI 생성 결과를 실제 문서로 제출하거나 법적 판단의 근거로 사용하기 전에 반드시 그 내용을 직접 확인하고, 필요한 경우 변호사 등 전문가의 검토를 받아야 합니다.
3. AI 생성 결과를 검토 없이 그대로 사용함에 따라 발생하는 불이익에 대해 운영자는 제9조(책임의 제한)에서 정하는 범위를 초과하여 책임지지 않습니다.

**제4조 (이용 자격)**
서비스는 만 14세 이상만 이용할 수 있습니다. 만 14세 미만인 경우 회원가입 및 서비스 이용이 제한됩니다.

**제5조 (계정 관리)**
1. 이용자는 회원가입 시 정확한 정보를 제공해야 하며, 이메일과 비밀번호 등 계정 정보를 스스로의 책임으로 관리해야 합니다.
2. 이용자는 자신의 계정을 제3자에게 공유·양도·대여할 수 없습니다.
3. 계정 정보의 도용이나 무단 사용을 인지한 경우 이용자는 지체 없이 운영자에게 통지해야 하며, 통지 전까지 발생한 불이익에 대해 운영자는 책임지지 않습니다. 다만 운영자의 고의 또는 중과실이 있는 경우는 제외합니다.

**제6조 (이용자의 의무와 금지행위)**
1. 이용자는 사건 정보를 입력할 때 사실에 부합하는 내용을 입력해야 하며, 상대방 등 제3자의 개인정보는 이용자가 적법하게 제공할 수 있는 범위 내에서 필요한 최소한만 입력해야 합니다.
2. 이용자는 다음 각 호의 행위를 해서는 안 됩니다.
   - 허위임을 알면서 사건 정보를 입력하거나 무고·협박 등 불법적인 목적으로 이용하는 행위
   - 타인을 사칭하거나 타인의 명예를 훼손·비방할 목적으로 서비스를 이용하는 행위
   - 법령을 위반하거나 공서양속에 반하는 목적으로 서비스를 이용하는 행위
   - 서비스를 무단으로 복제·역설계하거나 자동화된 수단(크롤링 등)으로 접근하는 행위
   - 서비스의 정상적인 운영을 방해하는 행위
3. 이용자가 위 의무를 위반해 발생한 결과에 대한 책임은 이용자 본인에게 있습니다.

**제7조 (서비스의 변경 및 중단)**
1. 서비스는 개인이 운영하는 베타 성격의 서비스로, 운영자는 서비스의 전부 또는 일부를 변경하거나 중단할 수 있습니다.
2. 운영자는 서비스 내용을 변경하거나 중단하는 경우 사전에 서비스 내 공지 등 합리적인 방법으로 이용자에게 알립니다. 다만 긴급한 장애 대응 등 부득이한 사유가 있는 경우 사후에 알릴 수 있습니다.

**제8조 (지식재산권)**
1. 이용자는 자신이 입력하는 정보에 대해 서비스에 제공할 적법한 권한을 보유해야 합니다.
2. 운영자는 서비스 제공(저장, AI 처리, 화면 표시 등) 목적의 범위 내에서만 이용자가 입력한 정보를 사용하며, 이를 초과하는 방식으로 이용하지 않습니다.
3. 서비스의 명칭, 디자인, 소프트웨어 등 서비스 자체에 관한 권리는 운영자 또는 정당한 권리자에게 있으며, 이용자는 운영자의 사전 동의 없이 이를 복제·배포·상업적으로 이용할 수 없습니다.

**제9조 (책임의 제한)**
1. 운영자는 천재지변, 불가항력, 이용자의 귀책사유 등 운영자가 통제할 수 없는 사유로 발생한 손해에 대해 책임지지 않습니다.
2. 운영자는 AI 생성 결과의 내용을 이용자가 검증 없이 그대로 사용함으로써 발생한 손해에 대해, 운영자에게 고의 또는 중과실이 없는 한 책임지지 않습니다.
3. 이 조는 운영자의 고의 또는 중과실로 인한 책임이나 관련 법령상 제한할 수 없는 책임까지 면제하는 것은 아닙니다.

**제10조 (약관의 변경)**
1. 운영자는 관련 법령을 위반하지 않는 범위에서 이 약관을 변경할 수 있습니다.
2. 약관이 변경되는 경우 운영자는 변경 내용, 시행일 및 변경 사유를 서비스 내 공지 등의 방법으로 시행일 이전에 안내합니다.
3. 관련 법령상 이용자의 별도 동의가 필요한 변경의 경우에는 이용자의 동의를 받습니다.

**제11조 (개인정보 보호)**
운영자는 관련 법령에 따라 이용자의 개인정보를 보호합니다. 개인정보의 수집 항목, 이용 목적, 보유기간, 제3자 제공 및 국외 이전에 관한 사항은 이 약관이 아닌 별도의 "개인정보처리방침"에 따릅니다.

**제12조 (준거법 및 분쟁 해결)**
1. 이 약관 및 서비스 이용과 관련하여 운영자와 이용자 사이에 발생한 분쟁에는 대한민국 법을 적용합니다.
2. 분쟁에 관한 소송이 제기되는 경우 관련 법령에 따른 관할법원에 제기합니다.

**제13조 (운영자·문의처·시행일)**
- 운영자: {PRIVACY_OPERATOR_NAME}
- 문의 이메일: {PRIVACY_CONTACT_EMAIL}
- 이 약관은 {TERMS_EFFECTIVE_DATE}부터 적용됩니다.
"""


@st.dialog("이용약관", width="large")
def render_terms_of_service_dialog():
    if not terms_config_ready():
        st.info("이용약관을 준비 중입니다. 잠시 후 다시 확인해 주세요.")
    else:
        st.markdown(build_terms_of_service_text())


def render_case_evidence_manager(case_id: str) -> None:
    with st.expander("증거자료 첨부·관리"):
        st.caption(
            "사진·PDF·녹음파일을 최대 20개까지 첨부할 수 있습니다. "
            "파일당 최대 10MB이며 본인만 열람할 수 있습니다."
        )
        st.warning(
            "주민등록번호, 계좌 비밀번호, 카드번호, 인증번호 등 불필요한 민감정보는 "
            "가린 뒤 첨부해 주세요."
        )
        uploads = st.file_uploader(
            "증거자료 선택",
            type=["jpg", "jpeg", "png", "webp", "heic", "heif", "pdf", "mp3", "m4a", "wav"],
            accept_multiple_files=True,
            key=f"evidence_upload_{case_id}",
        )
        consent = st.checkbox(
            "[필수] 선택한 파일의 비공개 저장과 사건 관리 목적 이용에 동의합니다.",
            key=f"evidence_consent_{case_id}",
        )
        if st.button("선택한 파일 첨부", key=f"evidence_upload_btn_{case_id}"):
            if not uploads:
                st.warning("첨부할 파일을 선택해 주세요.")
            elif not consent:
                st.warning("파일을 첨부하려면 저장·이용에 동의해 주세요.")
            else:
                uploaded_count = 0
                errors = []
                for uploaded_file in uploads:
                    try:
                        upload_case_evidence(case_id, uploaded_file)
                        uploaded_count += 1
                    except ValueError as exc:
                        errors.append(f"{uploaded_file.name}: {exc}")
                    except Exception as exc:
                        log_evidence_upload_error("upload", exc)
                        errors.append(f"{uploaded_file.name}: {map_evidence_upload_error(exc)}")
                if uploaded_count:
                    st.success(f"증거자료 {uploaded_count}개를 첨부했습니다.")
                for error in errors:
                    st.error(error)
                if uploaded_count:
                    st.rerun()

        try:
            files = list_case_evidence(case_id)
        except Exception as exc:
            log_evidence_upload_error("list", exc)
            st.error("증거자료 목록을 불러오지 못했습니다. 잠시 후 다시 시도해 주세요.")
            return

        if not files:
            st.caption("첨부된 증거자료가 없습니다.")
            return

        selected = st.selectbox(
            "첨부된 파일",
            files,
            format_func=lambda item: item["display_name"],
            key=f"evidence_select_{case_id}",
        )
        download_key = f"evidence_download_{case_id}"
        col_prepare, col_remove = st.columns(2)
        with col_prepare:
            if st.button("다운로드 준비", key=f"evidence_prepare_{case_id}"):
                try:
                    st.session_state[download_key] = {
                        "path": selected["path"],
                        "name": selected["display_name"],
                        "data": download_case_evidence(selected["path"]),
                    }
                except Exception as exc:
                    log_evidence_upload_error("download", exc)
                    st.error("파일을 불러오지 못했습니다.")
        with col_remove:
            if st.button("선택 파일 삭제", key=f"evidence_delete_{case_id}"):
                try:
                    delete_case_evidence_file(selected["path"])
                except Exception as exc:
                    log_evidence_upload_error("delete", exc)
                    st.error("파일 삭제에 실패했습니다.")
                else:
                    st.session_state.pop(download_key, None)
                    st.success("증거자료를 삭제했습니다.")
                    st.rerun()

        prepared = st.session_state.get(download_key)
        if prepared and prepared.get("path") == selected["path"]:
            st.download_button(
                "파일 다운로드",
                data=prepared["data"],
                file_name=prepared["name"],
                key=f"evidence_download_btn_{case_id}",
            )


def render_brand_header():
    st.markdown(
        '<div class="dgo-hero">'
        '<div class="dgo-hero-title">⚖️ 고소음이</div>'
        '<div class="dgo-hero-subtitle">어려운 피해 상황을 정리하고, 고소장 초안 작성을 돕는 AI 법률정보 도우미</div>'
        '</div>',
        unsafe_allow_html=True,
    )


def render_footer():
    st.markdown(
        '<div class="dgo-footer">'
        '고소음이는 법률정보와 문서 작성을 돕는 서비스이며, 변호사의 법률자문을 대체하지 않습니다.'
        '</div>',
        unsafe_allow_html=True,
    )


def render_signup_guide():
    st.markdown(
        "**가입 절차**\n\n"
        "① 이메일 입력 → ② 인증메일 확인(스팸함 포함) → ③ 앱으로 돌아와 로그인"
    )


GOOGLE_G_LOGO_SVG = (
    '<svg width="18" height="18" viewBox="0 0 48 48" aria-hidden="true">'
    '<path fill="#EA4335" d="M24 9.5c3.54 0 6.71 1.22 9.21 3.6l6.85-6.85C35.9 2.38 30.47 0 24 0 '
    '14.62 0 6.51 5.38 2.56 13.22l7.98 6.19C12.43 13.72 17.74 9.5 24 9.5z"/>'
    '<path fill="#4285F4" d="M46.98 24.55c0-1.57-.15-3.09-.38-4.55H24v9.02h12.94c-.58 2.96-2.26 '
    '5.48-4.78 7.18l7.73 6c4.51-4.18 7.09-10.36 7.09-17.65z"/>'
    '<path fill="#FBBC05" d="M10.53 28.59c-.48-1.45-.76-2.99-.76-4.59s.27-3.14.76-4.59l-7.98-6.19'
    'C.92 16.46 0 20.12 0 24c0 3.88.92 7.54 2.56 10.78l7.97-6.19z"/>'
    '<path fill="#34A853" d="M24 48c6.48 0 11.93-2.13 15.89-5.81l-7.73-6c-2.15 1.45-4.92 2.3-8.16 '
    '2.3-6.26 0-11.57-4.22-13.47-9.91l-7.98 6.19C6.51 42.62 14.62 48 24 48z"/>'
    '</svg>'
)

NAVER_N_LOGO_SVG = (
    '<svg width="18" height="18" viewBox="0 0 18 18" aria-hidden="true">'
    '<rect width="18" height="18" rx="2" fill="#03C75A"/>'
    '<path fill="#FFFFFF" d="M5 4h3l2.9 4.2V4H14v10h-3l-2.9-4.2V14H5V4z"/>'
    '</svg>'
)


def render_google_login_button(url: str) -> None:
    st.markdown(
        f'<a class="dgo-google-btn" href="{html.escape(url, quote=True)}" target="_blank" rel="noopener noreferrer">'
        f'{GOOGLE_G_LOGO_SVG}<span>Google로 계속하기</span></a>',
        unsafe_allow_html=True,
    )


def render_naver_login_button(url: str) -> None:
    st.markdown(
        f'<a class="dgo-google-btn" href="{html.escape(url, quote=True)}" target="_blank" rel="noopener noreferrer">'
        f'{NAVER_N_LOGO_SVG}<span>네이버로 로그인</span></a>',
        unsafe_allow_html=True,
    )


def render_social_login_section(key_suffix: str) -> None:
    google_attempted = google_login_available()
    google_url = get_google_login_url() if google_attempted else None
    show_naver = naver_login_available()
    log_auth_gate_debug(bool(google_url), google_attempted)
    if not google_url and not show_naver:
        return

    st.divider()
    st.caption(SOCIAL_LOGIN_NOTICE)
    if google_url:
        render_google_login_button(google_url)
    if show_naver:
        render_naver_login_button(NAVER_LOGIN_START_URL)


def render_google_age_gate():
    st.markdown(
        "**만 14세 이상 확인**\n\n"
        "고소음이는 만 14세 이상만 이용할 수 있습니다. "
        "Google 계정으로 처음 이용하시는 경우 아래 확인이 필요합니다."
    )
    google_age_consent = st.checkbox("[필수] 만 14세 이상입니다.", key="google_age_consent")
    if st.button("확인하고 계속하기", key="google_age_submit"):
        if not google_age_consent:
            st.error("만 14세 이상인 경우에만 서비스를 이용할 수 있습니다.")
        else:
            try:
                confirm_google_age()
                confirmed = True
            except Exception as error:
                log_auth_error(error)
                confirmed = False
            if confirmed:
                st.rerun()
            st.error("확인 저장 중 문제가 발생했습니다. 잠시 후 다시 시도해 주세요.")

    if st.button("로그아웃", key="google_age_logout_btn"):
        perform_logout()


def render_naver_age_gate(status_error: bool):
    if status_error:
        st.error("계정 확인 중 문제가 발생했습니다. 잠시 후 다시 시도해 주세요.")
    else:
        st.markdown(
            "**만 14세 이상 확인**\n\n"
            "고소음이는 만 14세 이상만 이용할 수 있습니다. "
            "네이버 계정으로 처음 이용하시는 경우 아래 확인이 필요합니다."
        )
        naver_age_consent = st.checkbox("[필수] 만 14세 이상입니다.", key="naver_age_consent")
        if st.button("확인하고 계속하기", key="naver_age_submit"):
            if not naver_age_consent:
                st.error("만 14세 이상인 경우에만 서비스를 이용할 수 있습니다.")
            else:
                try:
                    confirmed = get_authed_client().rpc("naver_confirm_age").execute().data is True
                except Exception:
                    confirmed = False
                if confirmed:
                    st.rerun()
                st.error("확인 저장 중 문제가 발생했습니다. 잠시 후 다시 시도해 주세요.")

    if st.button("로그아웃", key="naver_age_logout_btn"):
        perform_logout()


def render_pending_email_screen(email, resend_key, on_resend, message):
    st.success(message)
    st.info("받은 편지함에 메일이 보이지 않으면 **스팸함(정크메일함)**도 확인해 주세요.")

    elapsed = time.time() - (st.session_state.get(resend_key) or 0)
    remaining = max(0, 60 - int(elapsed))

    if remaining > 0:
        st.caption(f"{remaining}초 후 다시 보낼 수 있습니다. (화면을 조작하면 남은 시간이 갱신됩니다)")
        st.button("메일 다시 보내기", disabled=True, key=f"{resend_key}_disabled_btn")
    else:
        if st.button("메일 다시 보내기", key=f"{resend_key}_btn"):
            base_url = get_app_base_url()
            if not base_url:
                st.error("앱 주소 설정이 필요합니다. 관리자에게 문의해 주세요.")
            else:
                try:
                    on_resend(email, base_url)
                    st.session_state[resend_key] = time.time()
                    st.success("메일을 다시 보냈습니다.")
                except Exception as error:
                    log_auth_error(error)
                    st.error(map_auth_error(error))

    if st.button("로그인 화면으로", key=f"{resend_key}_back_btn"):
        st.session_state.auth_view = "login"
        st.rerun()


def render_check_email_screen(email, resend_key):
    def _resend(target_email, base_url):
        auth_client = get_supabase_anon_client()
        auth_client.auth.resend({
            "type": "signup",
            "email": target_email,
            "options": {"email_redirect_to": base_url},
        })

    render_pending_email_screen(
        email, resend_key, _resend,
        f"**{email}** 주소로 인증 메일을 보냈습니다. 메일을 확인해 주세요!",
    )
    st.info("인증 메일의 링크를 누른 다음, 이 앱으로 돌아와 로그인해 주세요.")


def render_reset_sent_screen():
    email = st.session_state.get("reset_sent_email") or ""

    def _resend(target_email, base_url):
        auth_client = get_supabase_anon_client()
        auth_client.auth.reset_password_for_email(target_email, {"redirect_to": base_url})

    render_pending_email_screen(
        email, "reset_resend_last_sent_at", _resend,
        f"**{email}** 주소로 비밀번호 재설정 메일을 보냈습니다.",
    )


def render_forgot_password_form():
    st.markdown("가입하신 이메일로 비밀번호 재설정 메일을 보내드립니다.")

    with st.form("forgot_password_form"):
        forgot_email = st.text_input("이메일", key="forgot_password_email")
        forgot_submitted = st.form_submit_button("재설정 메일 보내기")

    if forgot_submitted:
        base_url = get_app_base_url()
        if not base_url:
            st.error("앱 주소 설정이 필요합니다. 관리자에게 문의해 주세요.")
        else:
            try:
                auth_client = get_supabase_anon_client()
                auth_client.auth.reset_password_for_email(forgot_email, {"redirect_to": base_url})
                st.session_state.reset_sent_email = forgot_email
                st.session_state.reset_resend_last_sent_at = None
                st.session_state.auth_view = "reset_sent"
                st.rerun()
            except Exception as error:
                log_auth_error(error)
                st.error(map_auth_error(error))

    if st.button("로그인으로 돌아가기", key="forgot_password_back_btn"):
        st.session_state.auth_view = "login"
        st.rerun()


def render_recovery_set_password_form():
    st.markdown("새 비밀번호를 설정해 주세요.")

    with st.form("recovery_set_password_form"):
        new_password = st.text_input("새 비밀번호 (6자 이상)", type="password", key="recovery_new_password")
        new_password_confirm = st.text_input("새 비밀번호 확인", type="password", key="recovery_new_password_confirm")
        recovery_submitted = st.form_submit_button("비밀번호 변경")

    if recovery_submitted:
        if new_password != new_password_confirm:
            st.error("비밀번호가 일치하지 않습니다.")
        else:
            pending = st.session_state.get("pending_recovery_session")
            if not pending:
                st.session_state.auth_view = "login"
                set_auth_flash("error", "인증 세션이 만료되었습니다. 처음부터 다시 시도해 주세요.")
                st.rerun()
            try:
                client = get_session_authed_client(pending)
                client.auth.update_user({"password": new_password})
                refreshed = client.auth.get_session()
                if not refreshed:
                    raise RuntimeError("password update 이후 세션을 확인할 수 없습니다.")
                new_session = session_to_dict(refreshed)
                st.session_state.supabase_session = new_session
                st.session_state.pending_recovery_session = None
                st.session_state.auth_view = "login"
                set_auth_flash("success", "비밀번호가 변경되었습니다.")
                st.rerun()
            except Exception as error:
                log_auth_error(error)
                st.session_state.pending_recovery_session = None
                st.session_state.auth_view = "login"
                set_auth_flash("error", map_auth_error(error))
                st.rerun()

    if st.button("취소", key="recovery_cancel_btn"):
        st.session_state.pending_recovery_session = None
        st.session_state.auth_view = "login"
        st.rerun()


st.set_page_config(
    page_title="고소음이 | AI 법률정보 도우미",
    page_icon="⚖️",
    initial_sidebar_state="collapsed",
)
st.markdown(BRAND_CSS, unsafe_allow_html=True)

if "supabase_session" not in st.session_state:
    st.session_state.supabase_session = None
if "auth_view" not in st.session_state:
    st.session_state.auth_view = "login"
if "auth_flash" not in st.session_state:
    st.session_state.auth_flash = None
if "signup_sent_email" not in st.session_state:
    st.session_state.signup_sent_email = ""
if "signup_resend_last_sent_at" not in st.session_state:
    st.session_state.signup_resend_last_sent_at = None
if "reset_sent_email" not in st.session_state:
    st.session_state.reset_sent_email = ""
if "reset_resend_last_sent_at" not in st.session_state:
    st.session_state.reset_resend_last_sent_at = None
if "pending_recovery_session" not in st.session_state:
    st.session_state.pending_recovery_session = None

handle_auth_query_params()
handle_naver_query_params()
handle_google_query_params()
restore_browser_session()

if not st.session_state.supabase_session:
    render_brand_header()
    render_auth_flash()

    if st.session_state.auth_view == "recovery_set_password":
        render_recovery_set_password_form()
    elif st.session_state.auth_view == "signup_sent":
        render_check_email_screen(st.session_state.signup_sent_email, "signup_resend_last_sent_at")
    elif PASSWORD_RESET_ENABLED and st.session_state.auth_view == "reset_sent":
        render_reset_sent_screen()
    elif PASSWORD_RESET_ENABLED and st.session_state.auth_view == "forgot_password":
        render_forgot_password_form()
    else:
        if st.button("개인정보처리방침 보기", key="privacy_policy_view_btn_login"):
            render_privacy_policy_dialog()

        if st.button("이용약관 보기", key="terms_of_service_view_btn_login"):
            render_terms_of_service_dialog()

        login_tab, signup_tab = st.tabs(["로그인", "회원가입"])

        with login_tab:
            with st.form("login_form"):
                login_email = st.text_input("이메일", key="login_email")
                login_password = st.text_input("비밀번호", type="password", key="login_password")
                login_submitted = st.form_submit_button("로그인")

            if login_submitted:
                try:
                    auth_client = get_supabase_anon_client()
                    auth_response = auth_client.auth.sign_in_with_password(
                        {"email": login_email, "password": login_password}
                    )
                    if not getattr(auth_response, "session", None) or not getattr(auth_response, "user", None):
                        raise RuntimeError("empty sign_in response")
                    st.session_state.supabase_session = session_from_auth_response(auth_response)
                    st.rerun()
                except Exception as error:
                    log_auth_error(error)
                    st.error(map_auth_error(error))

            if PASSWORD_RESET_ENABLED:
                if st.button("비밀번호를 잊으셨나요?", key="forgot_password_link_btn"):
                    st.session_state.auth_view = "forgot_password"
                    st.rerun()

            render_social_login_section("login")

        with signup_tab:
            render_signup_guide()
            st.markdown(
                "**회원가입 시 수집하는 개인정보**\n"
                "- 수집 항목: 이메일, 인증을 위한 비밀번호\n"
                "- 목적: 회원가입·로그인·계정관리\n"
                "- 보관기간: 회원탈퇴 시까지\n"
                "- 동의를 거부할 수 있으나, 거부 시 회원가입이 제한됩니다.\n"
                "- 비밀번호는 Supabase Auth에서 암호화된 인증정보로 처리되며 운영자는 원문을 열람할 수 없습니다."
            )

            with st.form("signup_form"):
                signup_email = st.text_input("이메일", key="signup_email", max_chars=254)
                signup_password = st.text_input("비밀번호 (6자 이상)", type="password", key="signup_password", max_chars=128)
                signup_password_confirm = st.text_input("비밀번호 확인", type="password", key="signup_password_confirm", max_chars=128)
                signup_privacy_consent = st.checkbox(
                    "[필수] 위 안내를 확인했으며 개인정보 수집·이용에 동의합니다.",
                    key="signup_privacy_consent",
                )
                signup_terms_consent = st.checkbox(
                    "[필수] 이용약관에 동의합니다.",
                    key="signup_terms_consent",
                )
                signup_age_consent = st.checkbox("[필수] 만 14세 이상입니다.", key="signup_age_consent")
                signup_submitted = st.form_submit_button("회원가입")

            if signup_submitted:
                if not privacy_config_ready() or not terms_config_ready():
                    st.error(
                        "현재 서비스 운영 정보 설정이 진행 중이라 잠시 회원가입을 "
                        "이용하실 수 없습니다. 잠시 후 다시 시도해 주세요."
                    )
                elif not (signup_privacy_consent and signup_terms_consent and signup_age_consent):
                    st.error("회원가입을 진행하려면 이용약관 동의, 개인정보 수집·이용 동의, 만 14세 이상 확인에 모두 체크해 주세요.")
                elif len(signup_password) < 6:
                    st.error("비밀번호는 6자 이상 입력해 주세요.")
                elif signup_password != signup_password_confirm:
                    st.error("비밀번호가 일치하지 않습니다.")
                else:
                    base_url = get_app_base_url()
                    if not base_url:
                        st.error("앱 주소 설정이 필요합니다. 관리자에게 문의해 주세요.")
                    else:
                        try:
                            auth_client = get_supabase_anon_client()
                            auth_client.auth.sign_up({
                                "email": signup_email,
                                "password": signup_password,
                                "options": {
                                    "email_redirect_to": base_url,
                                    "data": {
                                        "terms_version": TERMS_EFFECTIVE_DATE,
                                        "privacy_version": PRIVACY_POLICY_EFFECTIVE_DATE,
                                        "terms_consent": True,
                                        "privacy_consent": True,
                                    },
                                },
                            })
                            st.session_state.signup_sent_email = signup_email
                            st.session_state.signup_resend_last_sent_at = None
                            st.session_state.auth_view = "signup_sent"
                            st.rerun()
                        except Exception as error:
                            log_auth_error(error)
                            st.error(map_auth_error(error))

            render_social_login_section("signup")

    render_footer()
    st.stop()

if ensure_active_session() is None:
    st.rerun()
render_auth_flash()

current_user_id = st.session_state.supabase_session["user_id"]
if st.session_state.get("case_workflow_user_id") != current_user_id:
    clear_user_session_data()
    st.session_state.case_workflow_user_id = current_user_id

if is_naver_session(st.session_state.supabase_session):
    try:
        naver_age_rows = get_authed_client().rpc("naver_age_status").execute().data or []
        naver_age_confirmed = bool(naver_age_rows and naver_age_rows[0].get("age_confirmed"))
    except Exception:
        naver_age_confirmed = None
    if not naver_age_confirmed:
        render_brand_header()
        render_naver_age_gate(status_error=naver_age_confirmed is None)
        render_footer()
        st.stop()
elif is_google_primary_session(st.session_state.supabase_session):
    if not social_age_confirmed(st.session_state.supabase_session):
        render_brand_header()
        render_google_age_gate()
        render_footer()
        st.stop()

if not has_current_consent(current_user_id):
    render_brand_header()
    st.markdown(build_terms_of_service_text())
    st.markdown(build_privacy_policy_text())
    reconsent_terms = st.checkbox("[필수] 이용약관에 동의합니다.", key="reconsent_terms")
    reconsent_privacy = st.checkbox("[필수] 개인정보 수집·이용에 동의합니다.", key="reconsent_privacy")
    if st.button("동의하고 계속하기", key="reconsent_submit"):
        if reconsent_terms and reconsent_privacy:
            try:
                get_authed_client().rpc(
                    "record_consent",
                    {
                        "p_terms_version": TERMS_EFFECTIVE_DATE,
                        "p_privacy_version": PRIVACY_POLICY_EFFECTIVE_DATE,
                    },
                ).execute()
                st.rerun()
            except Exception as error:
                error_code = re.sub(
                    r"[^A-Za-z0-9_.-]",
                    "",
                    str(getattr(error, "code", "") or type(error).__name__),
                )[:80] or "unknown"
                error_status = getattr(error, "status", None)
                print(f"[consent_error] code={error_code} status={error_status}")
                st.error(
                    "동의 저장 중 문제가 발생했습니다. "
                    f"관리자 확인용 오류 코드: {error_code}"
                )
        else:
            st.error("이용약관과 개인정보 수집·이용에 모두 동의해 주세요.")
    render_footer()
    st.stop()

with st.sidebar:
    st.markdown("**고소음이 데모**")
    st.caption(st.session_state.supabase_session.get("email") or "")
    if st.button("로그아웃", key="logout_btn_sidebar"):
        perform_logout()

    if st.button("개인정보처리방침 보기", key="privacy_policy_view_btn_sidebar"):
        render_privacy_policy_dialog()

    if st.button("이용약관 보기", key="terms_of_service_view_btn_sidebar"):
        render_terms_of_service_dialog()

    if PRIVACY_CONTACT_EMAIL:
        st.caption(f"회원탈퇴를 원하시면 {PRIVACY_CONTACT_EMAIL}로 가입하신 이메일을 포함해 요청해 주세요. 확인 후 계정과 저장된 사건을 삭제해 드립니다.")
    else:
        st.caption("회원탈퇴 절차를 준비 중입니다.")

render_brand_header()

with st.expander("계정 및 개인정보 메뉴"):
    if st.button("개인정보처리방침 보기", key="privacy_policy_view_btn_main"):
        render_privacy_policy_dialog()

    if st.button("이용약관 보기", key="terms_of_service_view_btn_main"):
        render_terms_of_service_dialog()

    if PRIVACY_CONTACT_EMAIL:
        st.caption(f"회원탈퇴를 원하시면 {PRIVACY_CONTACT_EMAIL}로 가입하신 이메일을 포함해 요청해 주세요. 확인 후 계정과 저장된 사건을 삭제해 드립니다.")
    else:
        st.caption("회원탈퇴 절차를 준비 중입니다.")

    if st.button("로그아웃", key="logout_btn_main"):
        perform_logout()

card_col1, card_col2, card_col3 = st.columns(3)
with card_col1:
    st.markdown('<div class="dgo-card">📝<b>피해 내용 정리</b></div>', unsafe_allow_html=True)
with card_col2:
    st.markdown('<div class="dgo-card">📄<b>고소장 초안 생성</b></div>', unsafe_allow_html=True)
with card_col3:
    st.markdown('<div class="dgo-card">🗂️<b>사건 저장 및 관리</b></div>', unsafe_allow_html=True)

tab_chat, tab_case, tab_case_list, tab_terms, tab_qa = st.tabs(["AI 상담", "고소장 작성", "내 사건", "법률 용어", "피해 유형별 Q&A"])

with tab_chat:
    st.caption("피해 상황을 편하게 이야기하면 핵심 사실을 정리할 수 있도록 도와드려요.")

    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    if prompt := st.chat_input("무엇을 도와드릴까요?", max_chars=3000):
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)

        try:
            answer = process_question(prompt, st.session_state.model_history)
        except Exception:
            answer = "답변 생성 중 문제가 발생했습니다. 잠시 후 다시 시도해 주세요."

        st.session_state.messages.append({"role": "assistant", "content": answer})
        with st.chat_message("assistant"):
            st.markdown(answer)

with tab_case:
    st.caption("사건 정보를 입력하면 고소장 초안을 자동으로 생성합니다.")
    st.markdown(
        "**사건 저장 시 수집하는 개인정보**\n"
        "- 저장 항목: 사건 유형, 신고인 정보, 상대방 정보, 사건 발생 일시, 사건 발생 장소, "
        "사건 경위, 피해 내용, 증거자료, 목격자 유무, 처벌 의사 (총 10개 항목)\n"
        "- 목적: 사건 저장·관리, AI 상담 및 고소장 초안 생성\n"
        "- 보관기간: 이용자가 직접 삭제하거나 회원탈퇴할 때까지\n"
        "- 동의를 거부할 수 있으나, 거부 시 사건 저장과 고소장 초안 생성이 제한됩니다."
    )
    st.warning(
        "입력하신 사건 정보(신고인·상대방 정보 등 개인정보 포함 가능)는 Supabase 데이터베이스에 저장되고, "
        "고소장 초안 생성 시 Google Gemini API로 전송되어 처리됩니다."
    )
    st.info(
        "상대방 등 제3자의 정보는 이용자가 적법하게 제공할 수 있는 범위에서 필요한 최소한만 입력해 주세요. "
        "주민등록번호, 계좌 비밀번호, 카드번호, 인증번호(OTP), 불필요한 주소·연락처는 입력하지 마세요."
    )

    if st.session_state.pending_draft_entry:
        st.warning(
            "사건은 저장되었습니다. 다만 고소장 초안 생성에 실패했습니다. "
            "새 사건을 제출하기 전에 초안을 다시 생성하거나 이 사건을 보류해 주세요."
        )
        col_retry, col_cancel = st.columns(2)
        with col_retry:
            if st.button("초안만 다시 생성", key="retry_draft_button"):
                entry = st.session_state.pending_draft_entry
                try:
                    st.session_state.case_draft = generate_case_draft(entry)
                    st.session_state.case_draft_id = entry["case_id"]
                    st.session_state.pending_draft_entry = None
                    st.rerun()
                except Exception:
                    st.error("고소장 초안 생성에 다시 실패했습니다. 잠시 후 다시 시도해 주세요.")
        with col_cancel:
            if st.button("이 사건 보류하고 새로 시작", key="cancel_pending_draft_button"):
                st.session_state.pending_draft_entry = None
                st.rerun()
    else:
        with st.form("case_form"):
            field_values = {}
            for label, key in CASE_FIELDS:
                if key in CASE_LONG_TEXT_FIELDS:
                    field_values[key] = st.text_area(label, max_chars=CASE_FIELD_MAX_CHARS[key])
                else:
                    field_values[key] = st.text_input(label, max_chars=CASE_FIELD_MAX_CHARS[key])

            case_data_consent = st.checkbox(
                "[필수] 위 안내를 확인했으며 사건 정보 저장·전송에 동의합니다.",
                key="case_data_consent",
            )
            submitted = st.form_submit_button("사건 저장 및 고소장 초안 생성")

        if submitted:
            if not privacy_config_ready():
                st.error("현재 서비스 운영 정보 설정이 진행 중이라 잠시 사건을 저장할 수 없습니다. 잠시 후 다시 시도해 주세요.")
            elif not case_data_consent:
                st.error("사건을 저장하려면 데이터 저장·전송 안내에 동의해 주세요.")
            else:
                missing_labels = [
                    label for label, key in CASE_FIELDS
                    if not field_values[key].strip()
                ]
                over_limit_errors = [
                    f"{label}은 최대 {CASE_FIELD_MAX_CHARS[key]}자까지 입력할 수 있습니다."
                    for label, key in CASE_FIELDS
                    if len(field_values[key]) > CASE_FIELD_MAX_CHARS[key]
                ]

                if missing_labels:
                    st.error("다음 항목을 입력해 주세요: " + ", ".join(missing_labels))
                elif over_limit_errors:
                    for message in over_limit_errors:
                        st.error(message)
                else:
                    user_id = st.session_state.supabase_session["user_id"]
                    fingerprint = compute_case_fingerprint(user_id, field_values)

                    is_duplicate = (
                        st.session_state.last_saved_fingerprint == fingerprint
                        and st.session_state.last_saved_at is not None
                        and (time.time() - st.session_state.last_saved_at) < DUPLICATE_SUBMIT_WINDOW_SECONDS
                    )

                    case_entry = None
                    if is_duplicate:
                        case_entry = st.session_state.last_saved_case_entry
                        st.info("방금 제출한 동일한 사건이 이미 저장되어 있어 다시 저장하지 않았습니다.")
                    else:
                        kst = datetime.timezone(datetime.timedelta(hours=9))
                        new_entry = {
                            "case_id": str(uuid.uuid4()),
                            "user_id": user_id,
                            "original_request": "Streamlit 고소장 작성 폼에서 제출됨",
                            **{key: field_values[key].strip() for _, key in CASE_FIELDS},
                            "created_at": datetime.datetime.now(kst).isoformat(),
                        }

                        try:
                            save_case(new_entry)
                        except Exception:
                            st.error("사건 저장 중 문제가 발생했습니다. 잠시 후 다시 시도해 주세요.")
                        else:
                            st.session_state.last_saved_case_entry = new_entry
                            st.session_state.last_saved_fingerprint = fingerprint
                            st.session_state.last_saved_at = time.time()
                            st.session_state.case_draft = None
                            st.session_state.case_draft_id = None
                            case_entry = new_entry

                    if case_entry is not None:
                        already_drafted = (
                            st.session_state.case_draft_id == case_entry["case_id"]
                            and st.session_state.case_draft
                        )
                        if not already_drafted:
                            try:
                                st.session_state.case_draft = generate_case_draft(case_entry)
                                st.session_state.case_draft_id = case_entry["case_id"]
                            except Exception:
                                st.session_state.pending_draft_entry = case_entry
                                st.rerun()

    if st.session_state.case_draft:
        st.subheader("생성된 고소장 초안")
        st.text_area("초안 내용", st.session_state.case_draft, height=400)

        col_txt_download, col_docx_download = st.columns(2)

        with col_txt_download:
            st.download_button(
                "TXT 다운로드",
                data=st.session_state.case_draft,
                file_name=f"case_draft_{st.session_state.case_draft_id}.txt",
                mime="text/plain",
            )

        with col_docx_download:
            docx_buffer = build_case_draft_docx(st.session_state.case_draft)
            st.download_button(
                "Word 다운로드",
                data=docx_buffer,
                file_name=f"case_draft_{st.session_state.case_draft_id}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

with tab_case_list:
    st.caption("저장된 사건을 조회하고 삭제할 수 있습니다.")
    st.caption("계정과 사건을 모두 삭제(회원탈퇴)하려면 사이드바의 안내를 참고해 주세요.")

    cases, load_error = load_cases_safe()

    if load_error:
        st.error(load_error)
    elif not cases:
        st.info("저장된 사건이 없습니다.")
    else:
        for case in cases:
            case_id = case.get("case_id", "")

            st.markdown(
                f"**사건 유형**: {case.get('case_type', '확인 필요')} | "
                f"**신고인**: {case.get('reporter_info', '확인 필요')} | "
                f"**상대방**: {case.get('suspect_info', '확인 필요')} | "
                f"**작성일시**: {case.get('created_at', '확인 필요')} | "
                f"**case_id**: {case_id}"
            )

            with st.expander("자세히 보기"):
                for label, key in CASE_FIELDS:
                    st.markdown(f"**{label}**: {case.get(key, '확인 필요')}")
                st.markdown(f"**작성일시**: {case.get('created_at', '확인 필요')}")
                st.markdown(f"**case_id**: {case_id}")

            if EVIDENCE_UPLOAD_ENABLED:
                render_case_evidence_manager(case_id)

            confirm_delete = st.checkbox("이 사건 삭제", key=f"confirm_delete_{case_id}")

            if st.button("삭제하기", key=f"delete_button_{case_id}"):
                if not confirm_delete:
                    st.warning("삭제 확인에 체크해 주세요.")
                else:
                    try:
                        delete_case(case_id)
                    except Exception:
                        st.error("사건 삭제 중 문제가 발생했습니다. 잠시 후 다시 시도해 주세요.")
                    else:
                        st.success(f"사건(case_id: {case_id})을 삭제했습니다.")
                        st.rerun()

            st.divider()

with tab_terms:
    st.subheader("어려운 법률 용어, 쉽게 보기")
    st.caption("형사 사건에서 자주 보이는 말의 일반적인 뜻을 설명합니다. 개별 사건의 결과를 판단하는 내용은 아닙니다.")
    term_query = st.text_input("용어 검색", placeholder="예: 고소, 피의자, 기소", key="legal_terms_query").strip().casefold()
    matching_terms = [
        (term, explanation) for term, explanation in LEGAL_TERMS
        if term_query in term.casefold() or term_query in explanation.casefold()
    ]
    if not matching_terms:
        st.info("검색 결과가 없습니다. 다른 단어로 검색해 주세요.")
    for term, explanation in matching_terms:
        with st.expander(term):
            st.write(explanation)
    st.markdown(f"[설명 근거: 대검찰청 사건처리절차 안내]({LEGAL_TERMS_SOURCE})")

with tab_qa:
    st.subheader("내 상황에 맞는 질문 살펴보기")
    st.caption("상황을 고르면 지금 해볼 일과 준비할 자료를 쉬운 말로 알려드립니다. 여기서 고르거나 읽은 내용은 사건에 저장되지 않습니다.")
    selected_qa_type = st.selectbox("피해 유형", list(VICTIM_QA), key="victim_qa_type")
    selected_qa = VICTIM_QA[selected_qa_type]
    for question, answer in selected_qa["questions"]:
        with st.expander(question):
            st.markdown(answer)
    st.markdown(f"[자세한 공식 안내 확인하기]({selected_qa['source']})")
    st.caption("위 내용은 일반 정보입니다. 실제 신고 방법과 보호조치는 담당 기관의 안내를 확인해 주세요.")

render_footer()
